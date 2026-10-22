"""Lectura del DOCX revisado y fusión de bancos en `QuestionBankManager` (BUG-7, BUG-10)."""

import os

import pandas as pd
from docx import Document

from pyexamgenerator.question_bank_manager import QuestionBankManager
from pyexamgenerator.question_generator import QuestionGenerator


def _write_docx(path, paragraphs):
    document = Document()
    document.add_heading("Banco de Preguntas - Pendiente de Revisar", level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return str(path)


def test_read_docx_keeps_colons_prefix_text_and_continuations(tmp_path):
    """BUG-7: los ':' cortaban el enunciado, 'A)' se borraba también dentro del texto y los
    párrafos de continuación se perdían."""
    docx_path = _write_docx(
        tmp_path / "revisadas.docx",
        [
            "Pregunta 1: Según la Ley 31/1995: ¿qué es la prevención?",
            "Tema: PRL - LPRL",
            "Estado: Aceptable",
            "A) Opción con A) dentro",
            "B) Dos",
            "C) Tres",
            "D) Cuatro",
            "Respuesta correcta: a",
            "Texto relevante: Primera línea.",
            "Segunda línea del texto relevante.",
            "",
            "Pregunta 2",
            "¿Enunciado escrito en el párrafo siguiente?",
            "Tema: PRL - RSP",
            "Estado: Pendiente de revisar",
            "A) Uno",
            "B) Dos",
            "C) Tres",
            "D) Cuatro",
            "Respuesta correcta: D",
            "Nota del revisor que no debe alterar la letra",
            "Texto relevante: ",
        ],
    )

    df = QuestionBankManager().read_questions_from_docx(docx_path)

    assert df is not None and len(df) == 2
    first, second = df.iloc[0], df.iloc[1]
    assert first["Número de pregunta"] == 1
    assert first["Pregunta"] == "Según la Ley 31/1995: ¿qué es la prevención?"
    assert first["Respuesta A"] == "Opción con A) dentro"
    assert first["Respuesta correcta"] == "A"
    assert first["Texto relevante"] == "Primera línea.\nSegunda línea del texto relevante."
    assert second["Número de pregunta"] == 2
    assert second["Pregunta"] == "¿Enunciado escrito en el párrafo siguiente?"
    assert second["Respuesta correcta"] == "D"
    assert second["Estado"] == "Pendiente de revisar"


def test_review_docx_round_trip(tmp_path):
    """El DOCX que escribe QuestionGenerator se vuelve a leer sin pérdidas."""
    original = pd.DataFrame(
        [
            {
                "Tema": "Tema 01",
                "Estado": "Pendiente de revisar",
                "Pregunta": "¿Cuál es la ratio: 3:1 o 1:3?",
                "Respuesta A": "3:1",
                "Respuesta B": "1:3",
                "Respuesta C": "Ambas (A) y (B)",
                "Respuesta D": "Ninguna",
                "Respuesta correcta": "A",
                "Texto relevante": "Se usa la ratio 3:1.",
            }
        ]
    )
    docx_path = str(tmp_path / "pendiente.docx")
    QuestionGenerator(api_key="dummy", model_name="models/test").save_questions_to_docx(original, docx_path)

    df = QuestionBankManager().read_questions_from_docx(docx_path)

    columns = ["Tema", "Estado", "Pregunta", "Respuesta A", "Respuesta B", "Respuesta C",
               "Respuesta D", "Respuesta correcta", "Texto relevante"]
    assert df[columns].to_dict("records") == original[columns].to_dict("records")


def test_add_reviewed_questions_returns_the_written_file(tmp_path, monkeypatch):
    """BUG-10: devolvía 'banco_de_preguntas_actualizado.xlsx' pero escribía otro nombre."""
    monkeypatch.chdir(tmp_path)
    row = {"Pregunta": "P1", "Respuesta A": "a", "Respuesta B": "b", "Respuesta C": "c", "Respuesta D": "d"}
    pd.DataFrame([row]).to_excel("banco.xlsx", index=False)
    pd.DataFrame([dict(row, Pregunta="P2")]).to_excel("revisadas.xlsx", index=False)

    saved_path = QuestionBankManager().add_reviewed_questions_to_existing_excel("banco.xlsx", "revisadas.xlsx")

    assert saved_path is not None and os.path.exists(saved_path)
    assert sorted(pd.read_excel(saved_path)["Pregunta"]) == ["P1", "P2"]
