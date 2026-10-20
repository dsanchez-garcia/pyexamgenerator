"""Regression tests for the review fixes in ExamGenerator and QuestionBankManager (QG-1 ... QG-22)."""
import xml.etree.ElementTree as ET

import pandas as pd
import pytest
from docx import Document

from pyexamgenerator.exam_generator import ExamGenerator
from pyexamgenerator.grading import SharedExamDataStore
from pyexamgenerator.question_bank_manager import QuestionBankManager


def _q(number, topic="Tema 01", estado="Aceptable", correct="A", statement=None, **extra):
    """One bank row; the option text of the correct letter is 'OK-<number>'."""
    row = {
        "Número de pregunta": number,
        "Tema": topic,
        "Estado": estado,
        "Pregunta": statement or f"Pregunta {number} de {topic}",
        "Respuesta A": f"A-{number}",
        "Respuesta B": f"B-{number}",
        "Respuesta C": f"C-{number}",
        "Respuesta D": f"D-{number}",
        "Respuesta correcta": correct,
        "Texto relevante": f"Relevante {number}",
    }
    letter = str(correct).strip().strip("().").upper() if isinstance(correct, str) else ""
    if letter in ("A", "B", "C", "D"):
        row[f"Respuesta {letter}"] = f"OK-{number}"
    row.update(extra)
    return row


def _write_bank(path, rows):
    pd.DataFrame(rows).to_excel(path, index=False)
    return str(path)


def _generate(bank_path, out_dir, generator=None, **kwargs):
    generator = generator or ExamGenerator()
    params = dict(
        bank_excel_path=str(bank_path),
        output_dir=str(out_dir),
        subject="Geo",
        exam="E",
        course="25-26",
        exam_names=["A"],
        usage_warning_enabled=False,
    )
    params.update(kwargs)
    generator.generate_exam_from_excel(**params)
    return generator


def _save_docx(path, lines):
    document = Document()
    document.add_heading("Banco de Preguntas - Pendiente de Revisar", level=1)
    for line in lines:
        document.add_paragraph(line)
    document.save(str(path))
    return str(path)


# --- QG-1: the bank is re-read on every generation ---------------------------------------------

def test_same_generator_uses_the_bank_passed_in_each_call(tmp_path):
    bank_a = _write_bank(tmp_path / "bank_a.xlsx", [_q(i, statement=f"BANCO A {i}") for i in range(1, 5)])
    bank_b = _write_bank(tmp_path / "bank_b.xlsx", [_q(i, "T9", statement=f"BANCO B {i}") for i in range(1, 7)])
    generator = ExamGenerator()

    _generate(bank_a, tmp_path, generator, exam="EA")
    _generate(bank_b, tmp_path, generator, exam="EB", update_excel=True)

    exam_b = pd.read_excel(tmp_path / "examen_Geo_EB_25-26_A_completo.xlsx")
    assert sorted(exam_b["Pregunta"]) == [f"BANCO B {i}" for i in range(1, 7)]
    bank_b_after = pd.read_excel(bank_b)
    assert bank_b_after["Pregunta"].tolist() == [f"BANCO B {i}" for i in range(1, 7)]
    assert bank_b_after["EB_25_26_uso"].tolist() == [1] * 6


def test_same_generator_sees_teacher_edits_and_does_not_revert_them(tmp_path):
    rows = [_q(i, estado="Aceptable" if i < 3 else "Pendiente de revisar") for i in range(1, 6)]
    bank = _write_bank(tmp_path / "bank.xlsx", rows)
    generator = _generate(bank, tmp_path, exam="E1")

    edited = pd.DataFrame(rows)
    edited["Estado"] = "Aceptable"
    edited.loc[0, "Pregunta"] = "Pregunta corregida"
    edited.to_excel(bank, index=False)
    _generate(bank, tmp_path, generator, exam="E2", update_excel=True)

    assert len(generator.exam_df) == 5
    after = pd.read_excel(bank)
    assert after.loc[0, "Pregunta"] == "Pregunta corregida"
    assert (after["Estado"] == "Aceptable").all()


# --- QG-2: update_excel marks the selected rows by index ---------------------------------------

def test_update_excel_marks_only_selected_rows_when_numbers_repeat(tmp_path):
    # A bank built from two generated files: both start numbering at 1.
    rows = [_q(i, "T1", statement=f"OLD {i}") for i in range(1, 5)]
    rows += [_q(i, "T2", estado="Pendiente de revisar", statement=f"NEW {i}") for i in range(1, 5)]
    bank = _write_bank(tmp_path / "bank.xlsx", rows)

    generator = _generate(bank, tmp_path, questions_per_topic={"T1": 2}, update_excel=True)

    after = pd.read_excel(bank)
    marked = sorted(after.loc[after["E_25_26_uso"] == 1, "Pregunta"])
    assert marked == sorted(generator.exam_df["Pregunta"])
    assert all(statement.startswith("OLD") for statement in marked)
    assert after.loc[after["Tema"] == "T2", "Veces usada en examen"].tolist() == [0, 0, 0, 0]


def test_update_excel_works_without_question_number_column(tmp_path):
    rows = [_q(i) for i in range(1, 4)]
    for row in rows:
        row.pop("Número de pregunta")
    bank = _write_bank(tmp_path / "bank.xlsx", rows)

    _generate(bank, tmp_path, update_excel=True)

    after = pd.read_excel(bank)
    assert after["E_25_26_uso"].tolist() == [1, 1, 1]
    assert after["Veces usada en examen"].tolist() == [1, 1, 1]


# --- QG-3: statements containing ':' (fix already in place, verified here) ---------------------

def test_read_questions_from_docx_keeps_colons_in_statement(tmp_path):
    docx_path = _save_docx(tmp_path / "rev.docx", [
        "Pregunta 1: Indique: según el art. 5: ¿cuál es el plazo?", "Tema: T1", "Estado: Aceptable",
        "A) a1", "B) b1", "C) c1", "D) d1", "Respuesta correcta: a", "Texto relevante: Art. 5: plazo",
        # A heading without ':' must not make the whole document unreadable.
        "Pregunta 2 ¿Sin dos puntos?", "Tema: T1", "Estado: Aceptable",
        "A) a2", "B) b2", "C) c2", "D) d2", "Respuesta correcta: b", "Texto relevante: x",
    ])

    df = QuestionBankManager().read_questions_from_docx(docx_path)

    assert df is not None and len(df) == 2
    assert df.loc[0, "Pregunta"] == "Indique: según el art. 5: ¿cuál es el plazo?"
    assert df.loc[0, "Texto relevante"] == "Art. 5: plazo"
    assert df.loc[1, "Respuesta B"] == "b2"


# --- QG-4: 'menos usadas' -------------------------------------------------------------------------

def test_least_used_selection_works_without_usage_column(tmp_path):
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(i) for i in range(1, 5)])

    generator = _generate(bank, tmp_path, questions_per_topic={"Tema 01": 2}, selection_method="menos usadas")

    # No usage recorded: every question counts as 0 uses and ties keep the bank order.
    assert generator.exam_df["Pregunta"].tolist() == ["Pregunta 1 de Tema 01", "Pregunta 2 de Tema 01"]


def test_least_used_selection_picks_never_used_questions_first(tmp_path):
    rows = [_q(i, statement=f"USADA {i}", P_24_25_uso=1, **{"Veces usada en examen": 1}) for i in range(1, 4)]
    rows += [_q(i, statement=f"NUEVA {i}") for i in range(10, 12)]  # added later: empty usage cells
    bank = _write_bank(tmp_path / "bank.xlsx", rows)

    generator = _generate(bank, tmp_path, questions_per_topic={"Tema 01": 2}, selection_method="menos usadas")

    assert sorted(generator.exam_df["Pregunta"]) == ["NUEVA 10", "NUEVA 11"]


# --- QG-5: 'Respuesta correcta' normalization and validation -----------------------------------

def test_correct_letter_variants_are_normalized(tmp_path):
    rows = [_q(1, correct="B "), _q(2, correct="b)"), _q(3, correct="(c)"), _q(4, correct="D."), _q(5, correct="a")]
    bank = _write_bank(tmp_path / "bank.xlsx", rows)

    _generate(bank, tmp_path, exam_names=["A", "B"], export_moodle_xml=True)

    for exam_type in ("A", "B"):
        exam = pd.read_excel(tmp_path / f"examen_Geo_E_25-26_{exam_type}_completo.xlsx")
        for _, row in exam.iterrows():
            correct_text = row[f"Respuesta {row['Respuesta correcta'].upper()}"]
            assert correct_text.startswith("OK-")


@pytest.mark.parametrize(
    "bad_row, expected_text",
    [
        (_q(2, correct=None), "Respuesta correcta"),
        (_q(2, correct="E"), "Respuesta correcta"),
        (_q(2, correct="C", **{"Respuesta C": None}), "Respuesta C"),
        (_q(2, **{"Pregunta": "  "}), "enunciado"),
    ],
)
def test_incomplete_selected_question_raises_clear_error(tmp_path, bad_row, expected_text):
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(1), bad_row])

    with pytest.raises(ValueError) as excinfo:
        _generate(bank, tmp_path)

    message = str(excinfo.value)
    assert expected_text in message
    assert "fila 3 del Excel" in message  # header is row 1, the faulty question is row 3
    assert not list(tmp_path.glob("examen_*"))


# --- QG-6 / QG-7: text written to Moodle XML ------------------------------------------------------

def test_moodle_xml_full_text_accepts_numeric_and_empty_cells(tmp_path):
    df = pd.DataFrame([_q(1, **{"Respuesta A": 1995, "Respuesta B": 2000.5, "Respuesta C": None,
                                 "Respuesta D": "2005"})])
    xml_path = tmp_path / "numeric.xml"

    ExamGenerator().generate_moodle_xml(df, str(xml_path), 1, exam="E", exam_type="1A", xml_use_answer_text=True)

    answers = [answer.findtext("text") for answer in ET.parse(str(xml_path)).getroot().iter("answer")]
    assert answers[0] == "1995"
    assert answers[1] == "2000.5"
    assert not answers[2]
    assert answers[3] == "2005"


def test_generate_exam_with_numeric_answers_exports_full_text_xml(tmp_path):
    rows = [_q(1), _q(2, correct="C", **{"Respuesta A": 1995, "Respuesta B": 1997, "Respuesta C": 2000,
                                         "Respuesta D": 2005})]
    bank = _write_bank(tmp_path / "bank.xlsx", rows)

    _generate(bank, tmp_path, export_moodle_xml=True, xml_use_answer_text=True)

    xml_text = (tmp_path / "examen_Geo_E_25-26_A.xml").read_text(encoding="utf-8")
    assert "1995" in xml_text and "2000" in xml_text


def test_moodle_xml_full_text_is_html_escaped(tmp_path):
    df = pd.DataFrame([_q(
        1,
        statement="Si x<y y z>w, ¿qué etiqueta <b> se usa & por qué?\nSegunda línea",
        **{"Respuesta A": "<p>", "Respuesta B": "a < b", "Respuesta C": "R&D", "Respuesta D": "ninguna"},
    )])
    xml_path = tmp_path / "html.xml"

    ExamGenerator().generate_moodle_xml(df, str(xml_path), 1, exam="E", exam_type="1A", xml_use_answer_text=True)

    root = ET.parse(str(xml_path)).getroot()
    question = root.find("./question[@type='multichoice']")
    # After XML parsing Moodle renders these values as HTML: markup characters must arrive escaped.
    statement = question.findtext("./questiontext/text")
    assert statement == "Si x&lt;y y z&gt;w, ¿qué etiqueta &lt;b&gt; se usa &amp; por qué?<br>Segunda línea"
    answers = [answer.findtext("text") for answer in question.findall("answer")]
    assert answers == ["&lt;p&gt;", "a &lt; b", "R&amp;D", "ninguna"]


def test_moodle_xml_letter_mode_still_builds_the_grading_answer_key(tmp_path):
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(i, correct="B") for i in range(1, 4)])

    _generate(bank, tmp_path, exam_names=["1A"], export_moodle_xml=True, penalty=-25)

    exam = pd.read_excel(tmp_path / "examen_Geo_E_25-26_1A_completo.xlsx")
    key = SharedExamDataStore().load_questions_xml(str(tmp_path / "examen_Geo_E_25-26_1A.xml"))
    assert len(key) == 3
    for _, row in exam.iterrows():
        answers = key[f"Pregunta {int(row['Número de pregunta']):02d}"]
        assert sorted(answers) == ["a", "b", "c", "d"]
        assert answers[row["Respuesta correcta"]] == pytest.approx(1.0)
        assert sorted(answers.values()) == pytest.approx([-0.25, -0.25, -0.25, 1.0])


# --- QG-12: file names ------------------------------------------------------------------------------

def test_output_file_names_replace_invalid_characters(tmp_path):
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(1), _q(2)])

    generator = _generate(bank, tmp_path / "out", subject="Física I/II", exam='Parcial 1: "T1"?',
                          exam_names=["1A"], export_moodle_xml=True)

    base = tmp_path / "out" / "examen_Física I_II_Parcial 1_ _T1___25-26_1A"
    for suffix in (".docx", "_completo.docx", "_completo.xlsx", ".xml"):
        assert base.with_name(base.name + suffix).exists()
    assert generator.generated_exams[0]["docx"] == str(base.with_name(base.name + ".docx"))
    # Titles and headers keep the original text.
    header = Document(str(base.with_name(base.name + ".docx"))).sections[0].header.paragraphs[0].text
    assert "Asignatura: Física I/II - Examen: Parcial 1: \"T1\"?" in header


# --- QG-14 (+ prefix inside option text): DOCX review reader ------------------------------------

def test_read_questions_from_docx_appends_continuation_paragraphs(tmp_path):
    docx_path = _save_docx(tmp_path / "rev.docx", [
        "Pregunta 1: ¿Cuál es el plazo", "de preaviso según el ET?", "Tema: T1", "Estado: Aceptable",
        "A) 15 días", "B) 30 días", "naturales", "C) 1 mes", "D) Ninguna", "Respuesta correcta: a",
        "Texto relevante: Art. 49", "con un preaviso de quince días.", "",
        "Pregunta 2: ¿Otra?", "Tema: T1", "Estado: Aceptable", "A) a", "B) b", "C) c", "D) d",
        "Respuesta correcta: b", "Texto relevante: y",
    ])

    df = QuestionBankManager().read_questions_from_docx(docx_path)

    assert df.loc[0, "Pregunta"] == "¿Cuál es el plazo\nde preaviso según el ET?"
    assert df.loc[0, "Respuesta B"] == "30 días\nnaturales"
    assert df.loc[0, "Texto relevante"] == "Art. 49\ncon un preaviso de quince días."
    assert df.loc[1, "Pregunta"] == "¿Otra?"
    assert df.loc[1, "Texto relevante"] == "y"


def test_read_questions_from_docx_keeps_keyword_repeated_inside_text(tmp_path):
    docx_path = _save_docx(tmp_path / "rev.docx", [
        "Pregunta 1: ¿Cuál?", "Tema: T1", "Estado: Aceptable", "A) Solo la A) del art. 5", "B) b",
        "C) Las opciones A) y C)", "D) d", "Respuesta correcta: a", "Texto relevante: x",
    ])

    df = QuestionBankManager().read_questions_from_docx(docx_path)

    assert df.loc[0, "Respuesta A"] == "Solo la A) del art. 5"
    assert df.loc[0, "Respuesta C"] == "Las opciones A) y C)"


# --- QG-15 / QG-16: paths returned by QuestionBankManager helpers -------------------------------

def test_generate_excel_from_docx_accepts_bare_file_name(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _save_docx(tmp_path / "rev.docx", [
        "Pregunta 1: ¿P?", "Tema: T1", "Estado: Aceptable", "A) a", "B) b", "C) c", "D) d",
        "Respuesta correcta: a", "Texto relevante: x",
    ])

    saved = QuestionBankManager().generate_excel_from_docx("rev.docx", "rev.xlsx")

    assert saved == "rev.xlsx"
    assert pd.read_excel(tmp_path / "rev.xlsx")["Pregunta"].tolist() == ["¿P?"]


def test_add_reviewed_questions_returns_the_path_it_writes(tmp_path, monkeypatch):
    work_dir = tmp_path / "otra_carpeta"
    work_dir.mkdir()
    monkeypatch.chdir(work_dir)
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(1)])
    reviewed = _write_bank(tmp_path / "reviewed.xlsx", [_q(2, statement="Otra pregunta")])

    saved = QuestionBankManager().add_reviewed_questions_to_existing_excel(bank, reviewed)

    assert saved == str(tmp_path / "bank_actualizado.xlsx")
    assert len(pd.read_excel(saved)) == 2
    assert len(pd.read_excel(bank)) == 1  # the original bank is left untouched
    assert not list(work_dir.iterdir())


# --- QG-17: duplicate detection when adding questions ------------------------------------------

def test_add_questions_without_duplicates_normalizes_like_unify(tmp_path):
    bank = _write_bank(tmp_path / "bank.xlsx", [
        _q(1, statement="¿Qué es un EPI?"),
        _q(2, statement="¿Plazo?", **{"Respuesta A": "15", "Respuesta B": "30"}),
    ])
    same_answers = {"Respuesta A": "OK-1", "Respuesta B": "B-1", "Respuesta C": "C-1", "Respuesta D": "D-1"}
    reviewed = _write_bank(tmp_path / "reviewed.xlsx", [
        _q(1, statement="¿Qué es un EPI? "),                    # trailing space
        _q(2, statement="¿qué es  un EPI?", **same_answers),    # case and double space
        _q(3, statement="¿Plazo?", **{"Respuesta A": "30", "Respuesta B": " 15", "Respuesta C": "C-2",
                                     "Respuesta D": "D-2"}),    # same answers, other order
        _q(4, statement="Nueva de verdad"),
        _q(5, statement="nueva de verdad",                      # duplicate inside the reviewed file
           **{"Respuesta A": "OK-4", "Respuesta B": "B-4", "Respuesta C": "C-4", "Respuesta D": "D-4"}),
    ])
    manager = QuestionBankManager()

    added_statement, df_statement = manager.add_questions_without_duplicates(bank, reviewed, ["Pregunta"])
    added_full, df_full = manager.add_questions_without_duplicates(bank, reviewed)

    assert added_statement == 1
    assert added_full == 1
    assert df_full["Pregunta"].tolist() == ["¿Qué es un EPI?", "¿Plazo?", "Nueva de verdad"]
    _, stats = manager.unify_question_banks([bank, reviewed], save=False)
    assert stats["unified"] == len(df_full)


# --- QG-18: topic matching --------------------------------------------------------------------------

@pytest.mark.parametrize("with_empty_topic", [False, True])
def test_topics_typed_as_numbers_match_dictionary_keys(tmp_path, with_empty_topic):
    rows = [_q(i, 1 if i < 4 else 2) for i in range(1, 7)]
    if with_empty_topic:  # an empty 'Tema' cell makes pandas read the column as float (1.0, 2.0)
        rows.append(_q(9, None, estado="Pendiente de revisar"))
    bank = _write_bank(tmp_path / "bank.xlsx", rows)

    generator = _generate(bank, tmp_path, questions_per_topic={"1": 2, "2": 1})

    assert len(generator.exam_df) == 3
    assert sorted(ExamGenerator._topic_key(topic) for topic in generator.exam_df["Tema"]) == ["1", "1", "2"]


def test_equivalent_topic_keys_never_repeat_a_question(tmp_path):
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(i, 1) for i in range(1, 4)])

    generator = _generate(bank, tmp_path, questions_per_topic={1: 2, "1": 1})

    assert len(generator.exam_df) == 3
    assert generator.exam_df["Pregunta"].is_unique


def test_equitable_distribution_merges_equivalent_topics():
    bank = pd.DataFrame({"Tema": [1, "1", 1.0, "2", 2, None]})
    allocation = ExamGenerator._distribute_questions_equitably(bank, 5)
    assert allocation == {"1": 3, "2": 2}


# --- QG-19: empty selection --------------------------------------------------------------------------

@pytest.mark.parametrize("questions_per_topic", [{"Tema 01": 0}, {}])
def test_empty_selection_raises_clear_error(tmp_path, questions_per_topic):
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(i) for i in range(1, 4)])

    with pytest.raises(ValueError, match="ninguna pregunta"):
        _generate(bank, tmp_path, questions_per_topic=questions_per_topic)

    assert not list(tmp_path.glob("examen_*"))


def test_no_topic_dictionary_still_uses_every_acceptable_question(tmp_path):
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(i) for i in range(1, 4)] + [_q(4, estado="Pendiente")])

    generator = _generate(bank, tmp_path, questions_per_topic=None)

    assert len(generator.exam_df) == 3


# --- QG-20: text in usage columns / numeric answer columns (pandas 3) --------------------------

def test_update_bank_with_exam_with_text_in_usage_column(tmp_path):
    bank = _write_bank(tmp_path / "bank.xlsx",
                       [_q(i, statement=f"U{i}", Ex_uso="x" if i == 1 else None) for i in range(1, 4)])
    exam = _write_bank(tmp_path / "exam.xlsx", [_q(1, statement="U2")])

    matched, df_bank = QuestionBankManager().update_bank_with_exam(bank, exam, exam_label="Ex")

    assert matched == 1
    assert df_bank["Ex_uso"].tolist()[:2] == ["x", 1]  # the teacher's mark is kept
    assert df_bank["Veces usada en examen"].tolist() == [0, 1, 0]  # text does not count as a use


@pytest.mark.parametrize("previous_marks", [["x", None, None], [True, False, False]])
def test_update_excel_with_text_in_existing_usage_column(tmp_path, previous_marks):
    bank = _write_bank(tmp_path / "bank.xlsx",
                       [_q(i, E_25_26_uso=mark) for i, mark in enumerate(previous_marks, start=1)])

    _generate(bank, tmp_path, questions_per_topic={"Tema 01": 3}, update_excel=True)

    after = pd.read_excel(bank)
    assert after["E_25_26_uso"].tolist() == [1, 1, 1]
    assert after["Veces usada en examen"].tolist() == [1, 1, 1]


def test_generate_question_text_shuffles_numeric_answer_columns(tmp_path):
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(i, **{"Respuesta B": 10 * i}) for i in range(1, 4)])
    df = pd.read_excel(bank)
    assert pd.api.types.is_integer_dtype(df["Respuesta B"])

    _, full_text = ExamGenerator().generate_question_text(df.copy(), renumber=False, shuffle_answers=True)

    assert full_text.count("Respuesta correcta: ") == 3
    assert "Respuesta correcta: None" not in full_text


# --- QG-21: check=True preview --------------------------------------------------------------------

def test_check_preview_accepts_si_and_is_saved_next_to_the_exams(tmp_path, monkeypatch):
    work_dir = tmp_path / "cwd"
    work_dir.mkdir()
    monkeypatch.chdir(work_dir)
    answers = []
    monkeypatch.setattr("builtins.input", lambda prompt="": answers.append(prompt) or " Si ")
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(1), _q(2)])

    _generate(bank, tmp_path / "out", check=True)

    assert len(answers) == 1
    assert (tmp_path / "out" / "examen_Geo_E_25-26_A.docx").exists()
    assert not list((tmp_path / "out").glob("examen_previsualizacion_*"))  # removed after confirming
    assert not list(work_dir.iterdir())


def test_check_preview_rejected_cancels_generation(tmp_path, monkeypatch):
    monkeypatch.setattr("builtins.input", lambda prompt="": "no")
    bank = _write_bank(tmp_path / "bank.xlsx", [_q(1), _q(2)])

    generator = _generate(bank, tmp_path / "out", check=True)

    assert generator.generation_cancelled is True
    assert (tmp_path / "out" / "examen_previsualizacion_Geo_E_25-26_completo.docx").exists()
    assert not (tmp_path / "out" / "examen_Geo_E_25-26_A.docx").exists()


# --- QG-22: Moodle category ---------------------------------------------------------------------

@pytest.mark.parametrize("additional, expected", [
    (None, "$course$/top/Examen E - 1A"),
    ("", "$course$/top/Examen E - 1A"),
    ("Grupo 1", "$course$/top/Examen E - 1A - Grupo 1"),
])
def test_moodle_category_matches_suggestion(tmp_path, additional, expected):
    xml_path = tmp_path / "cat.xml"

    ExamGenerator().generate_moodle_xml(pd.DataFrame([_q(1)]), str(xml_path), 1, exam="E", exam_type="1A",
                                        xml_cat_additional_text=additional)

    category = ET.parse(str(xml_path)).getroot().findtext("./question[@type='category']/category/text")
    assert category == expected
    if not additional:
        assert category == ExamGenerator.suggest_moodle_config(exam="E", exam_type="1A")["category_name"]


# --- Bare relative bank path without output_dir (os.makedirs('')) ------------------------------

def test_bare_relative_bank_path_without_output_dir(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _write_bank(tmp_path / "bank.xlsx", [_q(1), _q(2)])

    ExamGenerator().generate_exam_from_excel(
        bank_excel_path="bank.xlsx", subject="Geo", exam="E", course="25-26", exam_names=["A"],
        usage_warning_enabled=False,
    )

    assert (tmp_path / "examen_Geo_E_25-26_A.docx").exists()
