import random

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from pyexamgenerator.exam_generator import ExamGenerator
from pyexamgenerator.question_bank_manager import QuestionBankManager

# Custom name for the data sheet (instead of pandas' default 'Sheet1') and an extra sheet the
# teacher keeps in the same workbook.
DATA_SHEET = "Banco"
EXTRA_SHEET = "Notas"
NOTES_TEXT = "No borrar: notas del profesor"
NOTES_FORMULA = f"=COUNTA({DATA_SHEET}!A:A)-1"


def _build_bank_df(num_per_topic):
    """Builds a small acceptable question bank."""
    rows = []
    number = 1
    for topic, count in num_per_topic.items():
        for i in range(count):
            rows.append(
                {
                    "Número de pregunta": number,
                    "Tema": topic,
                    "Estado": "aceptable",
                    "Pregunta": f"{topic} - Pregunta {i + 1}",
                    "Respuesta A": f"{topic}-{i}-A",
                    "Respuesta B": f"{topic}-{i}-B",
                    "Respuesta C": f"{topic}-{i}-C",
                    "Respuesta D": f"{topic}-{i}-D",
                    "Respuesta correcta": "a",
                    "Texto relevante": f"Relevante {topic} {i}",
                }
            )
            number += 1
    return pd.DataFrame(rows)


def _write_bank_with_extra_sheet(path, bank_df):
    """Writes the bank in a custom-named first sheet, followed by a formatted notes sheet."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        bank_df.to_excel(writer, sheet_name=DATA_SHEET, index=False)
        notes = writer.book.create_sheet(EXTRA_SHEET)
        notes["A1"] = NOTES_TEXT
        notes["A1"].font = Font(bold=True)
        notes["B1"] = NOTES_FORMULA


def _assert_extra_sheet_kept(path):
    """The data sheet keeps its name and position and the notes sheet is left untouched."""
    workbook = load_workbook(path)
    assert workbook.sheetnames == [DATA_SHEET, EXTRA_SHEET]
    notes = workbook[EXTRA_SHEET]
    assert notes["A1"].value == NOTES_TEXT
    assert notes["A1"].font.bold
    assert notes["B1"].value == NOTES_FORMULA


# --- Every write-back of the bank keeps the other sheets and the data sheet name ---------------

def test_generate_exam_update_excel_keeps_extra_sheets(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _write_bank_with_extra_sheet(bank_path, _build_bank_df({"Tema 01": 4}))

    generator = ExamGenerator()
    random.seed(20261020)
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="Hojas",
        course="25-26",
        exam_names=["A"],
        total_questions=2,
        export_moodle_xml=False,
        update_excel=True,
    )

    _assert_extra_sheet_kept(bank_path)
    bank = pd.read_excel(bank_path)
    assert len(bank) == 4
    assert bank["Hojas_25_26_uso"].sum() == 2
    assert bank["Veces usada en examen"].sum() == 2


def test_add_questions_and_save_keeps_extra_sheets(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    reviewed_path = tmp_path / "revisadas.xlsx"
    full_df = _build_bank_df({"Tema 01": 4})
    _write_bank_with_extra_sheet(bank_path, full_df.iloc[:2])
    full_df.to_excel(reviewed_path, index=False)

    manager = QuestionBankManager()
    added_count, df_updated = manager.add_questions_without_duplicates(str(bank_path), str(reviewed_path))
    assert added_count == 2
    assert manager.save_dataframe_to_excel(df_updated, str(bank_path), overwrite=True) == str(bank_path)

    _assert_extra_sheet_kept(bank_path)
    assert pd.read_excel(bank_path)["Pregunta"].tolist() == full_df["Pregunta"].tolist()


def test_update_bank_with_exam_and_save_keeps_extra_sheets(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    bank_df = _build_bank_df({"Tema 01": 4})
    _write_bank_with_extra_sheet(bank_path, bank_df)
    exam_path = tmp_path / "examen_completo.xlsx"
    bank_df.iloc[[0, 2]].to_excel(exam_path, index=False)

    manager = QuestionBankManager()
    matched, df_updated = manager.update_bank_with_exam(str(bank_path), str(exam_path), exam_label="Parcial1")
    assert matched == 2
    assert manager.save_dataframe_to_excel(df_updated, str(bank_path), overwrite=True) == str(bank_path)

    _assert_extra_sheet_kept(bank_path)
    assert pd.read_excel(bank_path)["Parcial1_uso"].tolist() == [1, 0, 1, 0]


def test_update_bank_with_exams_and_save_keeps_extra_sheets(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    bank_df = _build_bank_df({"Tema 01": 4})
    _write_bank_with_extra_sheet(bank_path, bank_df)
    exam1_path = tmp_path / "examen_parcial1_completo.xlsx"
    exam2_path = tmp_path / "examen_parcial2_completo.xlsx"
    bank_df.iloc[[0, 1]].to_excel(exam1_path, index=False)
    bank_df.iloc[[1, 3]].to_excel(exam2_path, index=False)

    manager = QuestionBankManager()
    total, df_updated, stats = manager.update_bank_with_exams(
        str(bank_path), [(str(exam1_path), "Parcial1"), (str(exam2_path), "Parcial2")]
    )
    assert total == 4
    assert manager.save_dataframe_to_excel(df_updated, str(bank_path), overwrite=True) == str(bank_path)

    _assert_extra_sheet_kept(bank_path)
    assert pd.read_excel(bank_path)["Veces usada en examen"].tolist() == [1, 2, 0, 1]


def test_unify_banks_into_existing_bank_keeps_extra_sheets(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    other_path = tmp_path / "otro.xlsx"
    full_df = _build_bank_df({"Tema 01": 3})
    _write_bank_with_extra_sheet(bank_path, full_df.iloc[:2])
    full_df.iloc[1:].to_excel(other_path, index=False)  # One duplicate and one new question.

    manager = QuestionBankManager()
    _, stats = manager.unify_question_banks([str(bank_path), str(other_path)], output_path=str(bank_path))
    assert stats["output_path"] == str(bank_path)
    assert stats["duplicates_removed"] == 1

    _assert_extra_sheet_kept(bank_path)
    assert pd.read_excel(bank_path)["Pregunta"].tolist() == full_df["Pregunta"].tolist()


# --- How the data sheet itself is written -----------------------------------------------------

def test_save_dataframe_replaces_whole_data_sheet(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _write_bank_with_extra_sheet(bank_path, _build_bank_df({"Tema 01": 5}))
    smaller_df = _build_bank_df({"Tema 02": 2})[["Tema", "Pregunta"]]

    manager = QuestionBankManager()
    assert manager.save_dataframe_to_excel(smaller_df, str(bank_path)) == str(bank_path)

    _assert_extra_sheet_kept(bank_path)
    # No rows or columns are left over from the previous, larger data.
    pd.testing.assert_frame_equal(pd.read_excel(bank_path), smaller_df)


def test_save_dataframe_to_new_file_keeps_single_default_sheet(tmp_path):
    new_path = tmp_path / "nuevo.xlsx"

    manager = QuestionBankManager()
    assert manager.save_dataframe_to_excel(_build_bank_df({"Tema 01": 2}), str(new_path)) == str(new_path)

    assert load_workbook(new_path).sheetnames == ["Sheet1"]
    assert len(pd.read_excel(new_path)) == 2


def test_save_dataframe_error_leaves_existing_bank_untouched(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _write_bank_with_extra_sheet(bank_path, _build_bank_df({"Tema 01": 3}))
    original_bytes = bank_path.read_bytes()

    # A control character cannot be stored in an XLSX cell, so writing the last row fails midway.
    bad_df = _build_bank_df({"Tema 01": 3})
    bad_df.loc[2, "Pregunta"] = "Pregunta con \x01 caracter de control"

    manager = QuestionBankManager()
    assert manager.save_dataframe_to_excel(bad_df, str(bank_path)) is None
    assert bank_path.read_bytes() == original_bytes
