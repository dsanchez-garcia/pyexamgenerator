"""Ficheros que produce `ExamGenerator`: nombres, categoría de Moodle y plantillas (BUG-11..13)."""

import xml.etree.ElementTree as ET

import pandas as pd
import pytest
from docx import Document

from pyexamgenerator.exam_generator import ExamGenerator


def _bank_path(tmp_path):
    rows = [
        {
            "Número de pregunta": i + 1,
            "Tema": "Tema 01",
            "Estado": "Aceptable",
            "Pregunta": f"Pregunta {i + 1}",
            "Respuesta A": "a",
            "Respuesta B": "b",
            "Respuesta C": "c",
            "Respuesta D": "d",
            "Respuesta correcta": "a",
            "Texto relevante": "",
        }
        for i in range(2)
    ]
    path = tmp_path / "banco.xlsx"
    pd.DataFrame(rows).to_excel(path, index=False)
    return str(path)


def _generate(tmp_path, out="out", **kwargs):
    params = dict(
        bank_excel_path=_bank_path(tmp_path),
        output_dir=str(tmp_path / out),
        subject="Geo",
        exam="Parcial",
        course="25-26",
        exam_names=["A"],
    )
    params.update(kwargs)
    ExamGenerator().generate_exam_from_excel(**params)
    return tmp_path / out


def test_file_names_are_safe_for_any_system(tmp_path):
    """BUG-11: '/', ':' o '?' en la asignatura o el examen rompían la ruta; None daba 'None'."""
    out = _generate(tmp_path, subject="Física I/II", exam="Parcial: ¿final?")

    names = sorted(p.name for p in out.iterdir())
    assert "examen_Física I-II_Parcial- ¿final-_25-26_A.docx" in names
    assert not any(ch in name for name in names for ch in '/\\:*?"<>|')


def test_missing_subject_does_not_write_none(tmp_path):
    out = _generate(tmp_path, subject=None)

    names = [p.name for p in out.iterdir()]
    assert "examen__Parcial_25-26_A.docx" in names
    assert not any("None" in name for name in names)


def _xml_category(xml_path):
    root = ET.parse(str(xml_path)).getroot()
    return root.find("./question[@type='category']/category/text").text


def test_xml_category_matches_the_suggested_name(tmp_path):
    """BUG-12: el XML terminaba en ' -' mientras que la sugerencia no."""
    out = _generate(tmp_path, export_moodle_xml=True)

    category = _xml_category(out / "examen_Geo_Parcial_25-26_A.xml")
    assert category == "$course$/top/Examen Parcial - A"
    assert category == ExamGenerator.suggest_moodle_config(exam="Parcial", exam_type="A")["category_name"]


def test_xml_category_with_additional_text(tmp_path):
    out = _generate(tmp_path, export_moodle_xml=True, xml_cat_additional_text="Grupo 1")

    assert _xml_category(out / "examen_Geo_Parcial_25-26_A.xml") == "$course$/top/Examen Parcial - A - Grupo 1"


def test_template_placeholders_split_across_runs_and_in_header(tmp_path):
    """BUG-13: los marcadores partidos entre 'runs' de Word (o en la cabecera) no se sustituían."""
    template = Document()
    paragraph = template.add_paragraph()
    paragraph.add_run("Asignatura: {{sub")
    paragraph.add_run("ject}}").bold = True
    paragraph.add_run(" ({{course}})")
    template.sections[0].header.paragraphs[0].text = "Universidad - {{exam}} - {{exam_type}}"
    template_path = tmp_path / "plantilla.docx"
    template.save(str(template_path))

    out = _generate(tmp_path, template_docx_path=str(template_path))

    for name in ("examen_Geo_Parcial_25-26_A.docx", "examen_Geo_Parcial_25-26_A_completo.docx"):
        rendered = Document(str(out / name))
        body_text = "\n".join(p.text for p in rendered.paragraphs)
        assert "Asignatura: Geo (25-26)" in body_text
        header_text = "\n".join(p.text for p in rendered.sections[0].header.paragraphs)
        # The template header is kept (not overwritten by the student identification line).
        assert "Universidad - Parcial - A" in header_text
        assert "{{" not in body_text + header_text

    student_header = Document(str(out / "examen_Geo_Parcial_25-26_A.docx")).sections[0].header
    assert "Asignatura: Geo - Examen: Parcial" in "\n".join(p.text for p in student_header.paragraphs)


def test_missing_template_raises_instead_of_using_a_blank_document(tmp_path):
    """BUG-13: una ruta de plantilla equivocada se ignoraba en silencio."""
    with pytest.raises(FileNotFoundError, match="plantilla"):
        _generate(tmp_path, template_docx_path=str(tmp_path / "no_existe.docx"))
    assert not (tmp_path / "out").exists() or not any((tmp_path / "out").iterdir())
