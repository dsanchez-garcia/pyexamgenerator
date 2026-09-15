import xml.etree.ElementTree as ET
from pathlib import Path

import pandas as pd
import pytest
from docx import Document

from pyexamgenerator.exam_generator import ExamGenerator
from pyexamgenerator.question_bank_manager import QuestionBankManager


def _build_minimal_bank_df():
    return pd.DataFrame(
        [
            {
                "Número de pregunta": 1,
                "Tema": "Tema 01",
                "Estado": "aceptable",
                "Pregunta": "¿Capital de Francia?",
                "Respuesta A": "Madrid",
                "Respuesta B": "Paris",
                "Respuesta C": "Lisboa",
                "Respuesta D": "Roma",
                "Respuesta correcta": "b",
                "Texto relevante": "Paris es la capital de Francia.",
            }
        ]
    )


def test_exam_generator_accepts_estado_case_insensitive(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_minimal_bank_df().to_excel(bank_path, index=False)

    generator = ExamGenerator()
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="Parcial",
        course="25-26",
        exam_names=["A"],
        export_moodle_xml=False,
        update_excel=False,
    )

    assert (tmp_path / "examen_Geo_Parcial_25-26_A.docx").exists()


def test_generate_moodle_xml_non_blank_uses_question_and_answer_text(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_minimal_bank_df().to_excel(bank_path, index=False)

    generator = ExamGenerator()
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="Parcial",
        course="25-26",
        exam_names=["A"],
        export_moodle_xml=True,
        xml_use_answer_text=True,
        update_excel=False,
    )

    xml_path = tmp_path / "examen_Geo_Parcial_25-26_A.xml"
    root = ET.parse(str(xml_path)).getroot()
    question = root.find("./question[@type='multichoice']")

    assert question is not None
    question_text = question.find("./questiontext/text")
    answers = [a.find("./text").text for a in question.findall("./answer")]

    assert question_text is not None
    assert "Capital de Francia" in question_text.text
    assert "Paris" in answers


def test_generate_moodle_xml_snaps_minus_33_to_moodle_supported_fraction(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_minimal_bank_df().to_excel(bank_path, index=False)

    generator = ExamGenerator()
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="Parcial",
        course="25-26",
        exam_names=["A"],
        export_moodle_xml=True,
        penalty=-33,
        update_excel=False,
    )

    xml_path = tmp_path / "examen_Geo_Parcial_25-26_A.xml"
    root = ET.parse(str(xml_path)).getroot()
    question = root.find("./question[@type='multichoice']")
    assert question is not None

    negative_fractions = [
        float(answer.get("fraction", "0"))
        for answer in question.findall("./answer")
        if float(answer.get("fraction", "0")) < 0
    ]
    assert len(negative_fractions) == 3
    assert negative_fractions == pytest.approx([-33.3333333, -33.3333333, -33.3333333], abs=1e-6)


def test_generate_moodle_xml_from_existing_exam_xlsx(tmp_path):
    exam_xlsx_path = tmp_path / "examen_Geo_Parcial_25-26_A_completo.xlsx"
    _build_minimal_bank_df().to_excel(exam_xlsx_path, index=False)

    generator = ExamGenerator()
    generated_xml = generator.generate_moodle_xml_from_existing_exam_xlsx(
        exam_xlsx_path=str(exam_xlsx_path),
        penalty=-33,
        xml_use_answer_text=True,
    )

    generated_xml_path = Path(generated_xml)
    assert generated_xml_path.exists()
    assert generated_xml_path.name == "examen_Geo_Parcial_25-26_A.xml"

    root = ET.parse(str(generated_xml_path)).getroot()
    category_text = root.find("./question[@type='category']/category/text")
    assert category_text is not None
    assert "Parcial" in category_text.text
    assert "A" in category_text.text

    question = root.find("./question[@type='multichoice']")
    assert question is not None
    question_text = question.find("./questiontext/text")
    assert question_text is not None
    assert "Capital de Francia" in question_text.text


def test_exam_generator_template_docx_replaces_placeholders(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_minimal_bank_df().to_excel(bank_path, index=False)

    template_path = tmp_path / "template.docx"
    template = Document()
    template.add_paragraph("Plantilla {{subject}} - {{course}}")
    template.save(str(template_path))

    generator = ExamGenerator()
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="Parcial",
        course="25-26",
        exam_names=["A"],
        template_docx_path=str(template_path),
        update_excel=False,
    )

    docx_path = tmp_path / "examen_Geo_Parcial_25-26_A.docx"
    rendered = Document(str(docx_path))
    full_text = "\n".join(p.text for p in rendered.paragraphs)
    assert "Plantilla Geo - 25-26" in full_text


def test_question_bank_manager_add_only_acceptable_is_case_insensitive(tmp_path):
    existing_path = tmp_path / "existing.xlsx"
    reviewed_path = tmp_path / "reviewed.xlsx"

    pd.DataFrame([], columns=["Pregunta", "Respuesta A", "Respuesta B", "Respuesta C", "Respuesta D", "Estado"]).to_excel(existing_path, index=False)
    pd.DataFrame(
        [
            {
                "Pregunta": "P1",
                "Respuesta A": "A",
                "Respuesta B": "B",
                "Respuesta C": "C",
                "Respuesta D": "D",
                "Estado": "aceptable",
            }
        ]
    ).to_excel(reviewed_path, index=False)

    manager = QuestionBankManager()
    added_count, updated_df = manager.add_questions_without_duplicates(
        existing_bank_path=str(existing_path),
        reviewed_questions_path=str(reviewed_path),
        add_only_acceptable=True,
    )

    assert added_count == 1
    assert updated_df is not None
    assert len(updated_df) == 1

