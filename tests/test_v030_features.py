import random

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

from pyexamgenerator.exam_generator import ExamGenerator
from pyexamgenerator.question_bank_manager import QuestionBankManager


def _build_bank_df(num_per_topic=None):
    """Builds a small acceptable question bank spread over three topics."""
    if num_per_topic is None:
        num_per_topic = {"Tema 01": 4, "Tema 02": 4, "Tema 03": 4}
    rows = []
    number = 1
    for topic, count in num_per_topic.items():
        for i in range(count):
            rows.append(
                {
                    "Número de pregunta": number,
                    "Tema": topic,
                    "Estado": "aceptable",
                    "Pregunta": f"{topic} - Pregunta {i + 1}",
                    "Respuesta A": f"{topic}-{i}-A",
                    "Respuesta B": f"{topic}-{i}-B",
                    "Respuesta C": f"{topic}-{i}-C",
                    "Respuesta D": f"{topic}-{i}-D",
                    "Respuesta correcta": "a",
                    "Texto relevante": f"Relevante {topic} {i}",
                }
            )
            number += 1
    return pd.DataFrame(rows)


# --- Task 1: total number of questions + distribution -----------------------------------------

def test_distribute_questions_equitably_even_split():
    bank = _build_bank_df({"Tema 01": 5, "Tema 02": 5, "Tema 03": 5})
    allocation = ExamGenerator._distribute_questions_equitably(bank, 6)
    assert allocation == {"Tema 01": 2, "Tema 02": 2, "Tema 03": 2}
    assert sum(allocation.values()) == 6


def test_distribute_questions_equitably_odd_total_gives_extra_to_first_topics():
    bank = _build_bank_df({"Tema 01": 5, "Tema 02": 5, "Tema 03": 5})
    allocation = ExamGenerator._distribute_questions_equitably(bank, 7)
    assert sum(allocation.values()) == 7
    # The first topic (alphabetical) receives the extra question.
    assert allocation["Tema 01"] == 3
    assert allocation["Tema 02"] == 2
    assert allocation["Tema 03"] == 2


def test_distribute_questions_equitably_redistributes_when_topic_is_short():
    # Tema 01 only has 1 question; its shortfall must spill over to the others.
    bank = _build_bank_df({"Tema 01": 1, "Tema 02": 5, "Tema 03": 5})
    allocation = ExamGenerator._distribute_questions_equitably(bank, 7)
    assert sum(allocation.values()) == 7
    assert allocation["Tema 01"] == 1
    assert allocation["Tema 02"] + allocation["Tema 03"] == 6


def test_generate_exam_with_total_questions_equitativo(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_bank_df({"Tema 01": 4, "Tema 02": 4, "Tema 03": 4}).to_excel(bank_path, index=False)

    generator = ExamGenerator()
    random.seed(20260601)
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="Total",
        course="25-26",
        exam_names=["A"],
        total_questions=9,
        total_distribution="equitativo",
        export_moodle_xml=False,
        update_excel=False,
    )

    exam_xlsx = tmp_path / "examen_Geo_Total_25-26_A_completo.xlsx"
    assert exam_xlsx.exists()
    df = pd.read_excel(exam_xlsx)
    assert len(df) == 9
    # 9 across 3 topics => 3 each.
    assert df["Tema"].value_counts().to_dict() == {"Tema 01": 3, "Tema 02": 3, "Tema 03": 3}


def test_generate_exam_with_total_questions_azar(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_bank_df({"Tema 01": 4, "Tema 02": 4, "Tema 03": 4}).to_excel(bank_path, index=False)

    generator = ExamGenerator()
    random.seed(20260601)
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="TotalAzar",
        course="25-26",
        exam_names=["A"],
        total_questions=5,
        total_distribution="azar",
        export_moodle_xml=False,
        update_excel=False,
    )

    exam_xlsx = tmp_path / "examen_Geo_TotalAzar_25-26_A_completo.xlsx"
    assert exam_xlsx.exists()
    df = pd.read_excel(exam_xlsx)
    assert len(df) == 5


# --- Task 3: answer sheet on an odd page + no blank line between questions ---------------------

def test_answer_sheet_starts_on_odd_page_section(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_bank_df({"Tema 01": 3, "Tema 02": 3}).to_excel(bank_path, index=False)

    generator = ExamGenerator()
    random.seed(20260601)
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="Layout",
        course="25-26",
        exam_names=["A"],
        export_moodle_xml=False,
        update_excel=False,
    )

    docx_path = tmp_path / "examen_Geo_Layout_25-26_A.docx"
    document = Document(str(docx_path))
    # The student docx must have a second section that starts on an odd page (the answer sheet).
    assert len(document.sections) >= 2
    assert document.sections[1].start_type == WD_SECTION.ODD_PAGE


def test_no_trailing_break_between_questions(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_bank_df({"Tema 01": 3}).to_excel(bank_path, index=False)

    generator = ExamGenerator()
    random.seed(20260601)
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="NoBlank",
        course="25-26",
        exam_names=["A"],
        export_moodle_xml=False,
        update_excel=False,
    )

    docx_path = tmp_path / "examen_Geo_NoBlank_25-26_A.docx"
    document = Document(str(docx_path))
    # Locate the question paragraphs (they start with "Pregunta ").
    question_paragraphs = [p for p in document.paragraphs if p.text.startswith("Pregunta ")]
    assert question_paragraphs, "No se encontraron párrafos de pregunta."
    for paragraph in question_paragraphs:
        last_run = paragraph.runs[-1]
        # The last run of a question must not end with a line break (no blank line).
        assert last_run._element.findall(qn("w:br")) == []
        assert not last_run.text.endswith("\n")


# --- Task 2: update bank usage from an existing exam ------------------------------------------

def test_update_bank_with_exam_marks_used_questions(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    bank_df = _build_bank_df({"Tema 01": 4, "Tema 02": 4})
    bank_df.to_excel(bank_path, index=False)

    # Build an "exam" file that uses a subset of the bank questions (answers shuffled / reordered).
    used = bank_df.iloc[[0, 2, 5]].copy()
    # Simulate the shuffle that exams apply to answers: swap A and B text. Matching is by statement.
    used = used.rename(columns={"Respuesta A": "Respuesta B", "Respuesta B": "Respuesta A"})
    exam_path = tmp_path / "examen_completo.xlsx"
    used.to_excel(exam_path, index=False)

    manager = QuestionBankManager()
    matched, df_updated = manager.update_bank_with_exam(
        str(bank_path), str(exam_path), exam_label="Parcial1 25-26"
    )

    assert matched == 3
    assert df_updated is not None
    usage_columns = [c for c in df_updated.columns if c.endswith("_uso")]
    assert usage_columns == ["Parcial1_25_26_uso"]
    assert df_updated["Parcial1_25_26_uso"].sum() == 3
    assert "Veces usada en examen" in df_updated.columns
    assert df_updated["Veces usada en examen"].sum() == 3
    # The matched statements are flagged.
    used_statements = set(used["Pregunta"])
    flagged = df_updated.loc[df_updated["Parcial1_25_26_uso"] == 1, "Pregunta"]
    assert set(flagged) == used_statements


def test_update_bank_with_exam_missing_file_returns_error(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_bank_df({"Tema 01": 2}).to_excel(bank_path, index=False)

    manager = QuestionBankManager()
    matched, df_updated = manager.update_bank_with_exam(
        str(bank_path), str(tmp_path / "no_existe.xlsx")
    )
    assert matched == -1
    assert df_updated is None


# --- Unify several question banks into one, with a selectable duplicate criterion --------------

def _q(statement, a, b, c, d, topic="Tema 01", estado="aceptable"):
    return {
        "Tema": topic,
        "Estado": estado,
        "Pregunta": statement,
        "Respuesta A": a,
        "Respuesta B": b,
        "Respuesta C": c,
        "Respuesta D": d,
        "Respuesta correcta": "a",
        "Texto relevante": "",
    }


def test_unify_banks_statement_only_dedups_by_statement(tmp_path):
    bank1 = pd.DataFrame([_q("¿Capital de Francia?", "París", "Londres", "Roma", "Berlín")])
    bank2 = pd.DataFrame([
        _q("¿Capital de Francia?", "Paris", "London", "Rome", "Madrid"),  # same statement, other answers
        _q("¿Capital de Italia?", "Roma", "París", "Madrid", "Berlín"),
    ])
    p1, p2, out = tmp_path / "b1.xlsx", tmp_path / "b2.xlsx", tmp_path / "unif.xlsx"
    bank1.to_excel(p1, index=False)
    bank2.to_excel(p2, index=False)

    manager = QuestionBankManager()
    df, stats = manager.unify_question_banks(
        [str(p1), str(p2)], str(out), duplicate_criterion="enunciado"
    )

    assert len(df) == 2
    assert stats["duplicates_removed"] == 1
    assert stats["total_read"] == 3
    assert set(df["Pregunta"]) == {"¿Capital de Francia?", "¿Capital de Italia?"}
    # The first bank wins: the kept "Francia" row keeps the first bank's answers.
    francia = df[df["Pregunta"] == "¿Capital de Francia?"].iloc[0]
    assert francia["Respuesta B"] == "Londres"
    assert out.exists()


def test_unify_banks_statement_and_answers_keeps_distinct_answers(tmp_path):
    bank1 = pd.DataFrame([_q("¿Capital de Francia?", "París", "Londres", "Roma", "Berlín")])
    bank2 = pd.DataFrame([_q("¿Capital de Francia?", "Paris", "London", "Rome", "Madrid")])
    p1, p2 = tmp_path / "b1.xlsx", tmp_path / "b2.xlsx"
    bank1.to_excel(p1, index=False)
    bank2.to_excel(p2, index=False)

    manager = QuestionBankManager()
    df, stats = manager.unify_question_banks(
        [str(p1), str(p2)], duplicate_criterion="enunciado_y_respuestas", save=False
    )
    # Same statement but different answers -> both kept.
    assert len(df) == 2
    assert stats["duplicates_removed"] == 0


def test_unify_banks_answer_set_is_order_independent(tmp_path):
    bank1 = pd.DataFrame([_q("Pregunta X", "A", "B", "C", "D")])
    bank2 = pd.DataFrame([_q("Pregunta X", "B", "A", "D", "C")])  # same options, reordered
    p1, p2 = tmp_path / "b1.xlsx", tmp_path / "b2.xlsx"
    bank1.to_excel(p1, index=False)
    bank2.to_excel(p2, index=False)

    manager = QuestionBankManager()
    df, stats = manager.unify_question_banks(
        [str(p1), str(p2)], duplicate_criterion="pregunta_respuestas", save=False
    )
    # Reordered identical options are treated as the same question.
    assert len(df) == 1
    assert stats["duplicates_removed"] == 1


def test_unify_three_banks_with_gui_alias_and_normalization(tmp_path):
    bank1 = pd.DataFrame([_q("Misma pregunta", "a", "b", "c", "d")])
    bank2 = pd.DataFrame([_q("  MISMA PREGUNTA ", "a", "b", "c", "d")])  # case/space variant
    bank3 = pd.DataFrame([_q("Otra pregunta", "a", "b", "c", "d")])
    paths = []
    for i, b in enumerate((bank1, bank2, bank3)):
        p = tmp_path / f"b{i}.xlsx"
        b.to_excel(p, index=False)
        paths.append(str(p))

    manager = QuestionBankManager()
    df, stats = manager.unify_question_banks(
        paths, duplicate_criterion="pregunta_unica", save=False
    )
    # The case/whitespace variant of "Misma pregunta" is a duplicate -> 2 unique remain.
    assert len(df) == 2
    assert stats["duplicates_removed"] == 1
    assert len(stats["banks"]) == 3


def test_unify_banks_requires_paths():
    manager = QuestionBankManager()
    try:
        manager.unify_question_banks([])
        assert False, "Debe lanzar ValueError con lista vacía"
    except ValueError:
        pass
