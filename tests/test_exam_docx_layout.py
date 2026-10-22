"""Maquetación del DOCX del examen: hoja de respuestas en página impar con columna NC, tabla de
datos del alumno, cabecera con huecos de identificación, estilos y marcadores de plantilla."""

import random

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

from pyexamgenerator.exam_generator import ExamGenerator


def _build_bank_df(num_per_topic=None):
    """Builds a small acceptable question bank spread over three topics."""
    if num_per_topic is None:
        num_per_topic = {"Tema 01": 4, "Tema 02": 4, "Tema 03": 4}
    rows = []
    number = 1
    for topic, count in num_per_topic.items():
        for i in range(count):
            rows.append(
                {
                    "Número de pregunta": number,
                    "Tema": topic,
                    "Estado": "aceptable",
                    "Pregunta": f"{topic} - Pregunta {i + 1}",
                    "Respuesta A": f"{topic}-{i}-A",
                    "Respuesta B": f"{topic}-{i}-B",
                    "Respuesta C": f"{topic}-{i}-C",
                    "Respuesta D": f"{topic}-{i}-D",
                    "Respuesta correcta": "a",
                    "Texto relevante": f"Relevante {topic} {i}",
                }
            )
            number += 1
    return pd.DataFrame(rows)


def _generate_student_docx(tmp_path, exam="Layout", num_questions=12):
    bank_path = tmp_path / "bank.xlsx"
    _build_bank_df({"Tema 01": num_questions}).to_excel(bank_path, index=False)
    generator = ExamGenerator()
    random.seed(20260601)
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam=exam,
        course="25-26",
        exam_names=["A"],
        export_moodle_xml=False,
        update_excel=False,
    )
    docx_path = tmp_path / f"examen_Geo_{exam}_25-26_A.docx"
    assert docx_path.exists()
    return Document(str(docx_path))


def _build_minimal_bank_df():
    return pd.DataFrame(
        [
            {
                "Número de pregunta": 1,
                "Tema": "Tema 01",
                "Estado": "aceptable",
                "Pregunta": "¿Capital de Francia?",
                "Respuesta A": "Madrid",
                "Respuesta B": "Paris",
                "Respuesta C": "Lisboa",
                "Respuesta D": "Roma",
                "Respuesta correcta": "b",
                "Texto relevante": "Paris es la capital de Francia.",
            }
        ]
    )


def test_answer_sheet_starts_on_odd_page_section(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_bank_df({"Tema 01": 3, "Tema 02": 3}).to_excel(bank_path, index=False)

    generator = ExamGenerator()
    random.seed(20260601)
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="Layout",
        course="25-26",
        exam_names=["A"],
        export_moodle_xml=False,
        update_excel=False,
    )

    docx_path = tmp_path / "examen_Geo_Layout_25-26_A.docx"
    document = Document(str(docx_path))
    # The student docx must have a second section that starts on an odd page (the answer sheet).
    assert len(document.sections) >= 2
    assert document.sections[1].start_type == WD_SECTION.ODD_PAGE


def test_no_trailing_break_between_questions(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_bank_df({"Tema 01": 3}).to_excel(bank_path, index=False)

    generator = ExamGenerator()
    random.seed(20260601)
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="NoBlank",
        course="25-26",
        exam_names=["A"],
        export_moodle_xml=False,
        update_excel=False,
    )

    docx_path = tmp_path / "examen_Geo_NoBlank_25-26_A.docx"
    document = Document(str(docx_path))
    # Locate the question paragraphs (they start with "Pregunta ").
    question_paragraphs = [p for p in document.paragraphs if p.text.startswith("Pregunta ")]
    assert question_paragraphs, "No se encontraron párrafos de pregunta."
    for paragraph in question_paragraphs:
        last_run = paragraph.runs[-1]
        # The last run of a question must not end with a line break (no blank line).
        assert last_run._element.findall(qn("w:br")) == []
        assert not last_run.text.endswith("\n")


def test_student_data_table_has_unequal_columns(tmp_path):
    document = _generate_student_docx(tmp_path, exam="DataCols")
    # The first table in the student docx is the 'Datos del alumno' table.
    student_table = document.tables[0]
    assert len(student_table.columns) == 2
    label_width = student_table.cell(0, 0).width
    value_width = student_table.cell(0, 1).width
    # The data column must be clearly wider than the label column.
    assert value_width is not None and label_width is not None
    assert value_width > label_width
    assert round(label_width.cm, 1) == 3.5


def test_answer_sheet_has_nc_column_centered(tmp_path):
    document = _generate_student_docx(tmp_path, exam="NCcol", num_questions=12)
    # The answer sheet is the second table.
    answer_table = document.tables[1]
    headers = [answer_table.cell(0, j).text for j in range(len(answer_table.columns))]
    assert headers == ["Pregunta", "a", "b", "c", "d", "NC"]

    # The whole table is centered on the page.
    assert answer_table.alignment == WD_TABLE_ALIGNMENT.CENTER

    # The answer option columns (a/b/c/d/NC) are ~1 cm wide; 'Pregunta' is wider (fits its content).
    for j in range(1, 6):
        assert round(answer_table.cell(0, j).width.cm, 1) == 1.0
    assert answer_table.cell(0, 0).width.cm > answer_table.cell(0, 1).width.cm

    # Cell content is centered (check the header and a data cell).
    assert answer_table.cell(0, 0).paragraphs[0].alignment == WD_PARAGRAPH_ALIGNMENT.CENTER
    assert answer_table.cell(1, 0).paragraphs[0].alignment == WD_PARAGRAPH_ALIGNMENT.CENTER


def test_student_docx_has_identification_blanks_in_repeated_header(tmp_path):
    document = _generate_student_docx(tmp_path, exam="HeaderFields", num_questions=12)
    header = document.sections[0].header
    assert header.tables, "La cabecera del alumno debe incluir una tabla de identificacion."
    header_cells_text = "\n".join(cell.text for table in header.tables for row in table.rows for cell in row.cells)
    for label in ["Nombre", "Apellidos", "DNI/NIE", "Firma"]:
        assert label in header_cells_text
    assert len(document.sections) >= 2
    assert document.sections[1].header.is_linked_to_previous


def test_question_statements_use_heading_2_for_word_navigation(tmp_path):
    document = _generate_student_docx(tmp_path, exam="Heading2", num_questions=4)
    question_paragraphs = [p for p in document.paragraphs if p.text.startswith("Pregunta ")]
    assert question_paragraphs
    for paragraph in question_paragraphs:
        assert paragraph.style.name == "Heading 2"
        assert all(run.font.bold for run in paragraph.runs if run.text)
        assert not paragraph.text.endswith("\n")


def test_exam_generator_template_docx_replaces_placeholders(tmp_path):
    bank_path = tmp_path / "bank.xlsx"
    _build_minimal_bank_df().to_excel(bank_path, index=False)

    template_path = tmp_path / "template.docx"
    template = Document()
    template.add_paragraph("Plantilla {{subject}} - {{course}}")
    template.save(str(template_path))

    generator = ExamGenerator()
    generator.generate_exam_from_excel(
        bank_excel_path=str(bank_path),
        output_dir=str(tmp_path),
        subject="Geo",
        exam="Parcial",
        course="25-26",
        exam_names=["A"],
        template_docx_path=str(template_path),
        update_excel=False,
    )

    docx_path = tmp_path / "examen_Geo_Parcial_25-26_A.docx"
    rendered = Document(str(docx_path))
    full_text = "\n".join(p.text for p in rendered.paragraphs)
    assert "Plantilla Geo - 25-26" in full_text
