# pyexamgenerator: A tool for generating exams from PDF files using AI.
# Copyright (C) 2024 Daniel Sánchez-García

# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.

# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.

# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

import math
import os
import re
import textwrap
import time
import traceback
import unicodedata
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple

import pandas as pd
from docx import Document

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - older installations
    from PyPDF2 import PdfReader

try:
    import pymupdf as fitz  # PyMuPDF >= 1.24.3
except ImportError:
    try:
        import fitz  # PyMuPDF (legacy module name)
    except ImportError:
        fitz = None

from .ai_providers import (
    AIProviderError,
    BaseAIClient,
    InvalidAPIKeyError,
    InvalidRequestError,
    ModelNotFoundError,
    ProviderNotAvailableError,
    QuotaExceededError,
    ServiceOverloadedError,
    create_ai_client,
    resolve_provider,
)
from .question_parsing import OPTION_LETTERS, parse_questions_response, question_identity

# QuotaExceededError and ServiceOverloadedError are re-exported here for backward compatibility.
__all__ = [
    "QuestionGenerator",
    "QuotaExceededError",
    "ServiceOverloadedError",
    "QUESTION_COLUMNS",
    "INPUT_MODES",
]

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
#: Image formats that not every provider accepts: they are converted to PNG before being sent.
_IMAGE_EXTENSIONS_TO_CONVERT = {".bmp", ".gif", ".tif", ".tiff"}

#: Columns of the generated question files (same layout as the question bank).
QUESTION_COLUMNS = [
    "Tema", "Estado", "Pregunta", "Respuesta A", "Respuesta B", "Respuesta C", "Respuesta D",
    "Respuesta correcta", "Texto relevante",
]

#: Supported ways of sending the source material to the model.
#: 'text': the PDF text is extracted and included in the prompt.
#: 'image': PDF pages are rendered as images (or image files are used) and attached.
#: 'pdf': the PDF itself (or the PDF fragment) is attached as a document.
INPUT_MODES = ("text", "image", "pdf")
CHUNKING_MODES = ("pages", "text_length")

#: Maximum characters per cell accepted by Excel.
_EXCEL_CELL_LIMIT = 32000

_FATAL_PROVIDER_ERRORS = (
    QuotaExceededError, ServiceOverloadedError, InvalidAPIKeyError, ModelNotFoundError, ProviderNotAvailableError,
    InvalidRequestError,
)

_FORMAT_INSTRUCTIONS = """1. **Incluye la respuesta correcta:** Para cada pregunta, indica la letra de la opción correcta (A, B, C o D), sin paréntesis.
2. **Proporciona el texto relevante:** Incluye un fragmento literal del material donde se encuentra la justificación de la respuesta correcta.
3. **Formato JSON:** Presenta las preguntas como un array JSON en el que cada elemento es un objeto con exactamente estas claves:
   * "Pregunta": el enunciado de la pregunta.
   * "Opciones": una lista con exactamente 4 opciones (A, B, C y D, en ese orden), sin la letra inicial ni paréntesis.
   * "Respuesta correcta": la letra de la opción correcta ("A", "B", "C" o "D").
   * "Texto relevante": el fragmento del material que justifica la respuesta correcta.
   Ejemplo del formato (con una sola pregunta):
   [{"Pregunta": "¿...?", "Opciones": ["...", "...", "...", "..."], "Respuesta correcta": "B", "Texto relevante": "..."}]
Puedes incluir preguntas del tipo: señala la afirmación incorrecta entre las siguientes opciones; o señala la opción correcta entre las siguientes.
Además, puedes incluir respuestas como todas son correctas, o todas son incorrectas. No uses opciones que dependan de su posición (como "todas las anteriores" o "ninguna de las anteriores"), porque el orden de las opciones se baraja en los exámenes. Las 4 opciones deben ser distintas entre sí."""

_OUTPUT_RULE_API = (
    "**Responde ÚNICAMENTE con el array JSON**, sin texto adicional antes ni después. Si la petición "
    "incluye un esquema JSON para la respuesta, sigue ese esquema (con sus nombres de clave)."
)

_FINAL_REMINDER_API = "Recuerda: responde únicamente con el JSON de las preguntas."
_FINAL_REMINDER_WEB_CHAT = "Recuerda: responde únicamente con un bloque de código ```json con el array de preguntas."

_OUTPUT_RULE_WEB_CHAT = (
    "**Responde ÚNICAMENTE con un bloque de código ```json que contenga el array JSON completo**, "
    "sin texto adicional antes ni después. Si la respuesta es demasiado larga, termínala en una "
    "pregunta completa: te pediré que continúes con las restantes en otro bloque."
)


class QuestionGenerator:
    """
    Generates multiple-choice questions from PDF files (or images) using an AI model.

    Three providers are supported through their APIs: Google Gemini (keys from Google
    AI Studio), OpenAI (ChatGPT models) and Anthropic (Claude models). Without an API key
    the class can still prepare the prompts to paste in the web chats (Google AI Studio,
    ChatGPT, Claude) with :meth:`build_generation_prompts`/:meth:`export_prompts` and import
    the answers copied from the chat with :meth:`import_questions_from_responses`.
    """

    def __init__(
            self,
            api_key: Optional[str] = None,
            model_name: Optional[str] = None,
            provider: Optional[str] = "gemini",
            ai_client: Optional[BaseAIClient] = None,
            max_output_tokens: Optional[int] = None,
            temperature: Optional[float] = None,
    ):
        """
        Initializes the QuestionGenerator class.

        Args:
            api_key (Optional[str]): API key of the provider. Not needed to build prompts or to
                import answers copied from a web chat. If empty, the environment variable of the
                provider is used (GEMINI_API_KEY, OPENAI_API_KEY or ANTHROPIC_API_KEY).
            model_name (Optional[str]): Model to use (e.g. 'gemini-flash-latest', 'gpt-6.1-sol',
                'claude-sonnet-5-5'). Required only to call the API.
            provider (Optional[str]): 'gemini' (default), 'openai' or 'anthropic'. Aliases such
                as 'google', 'chatgpt' or 'claude' are accepted.
            ai_client (Optional[BaseAIClient]): Preconfigured client (overrides the previous
                arguments; useful for tests).
            max_output_tokens (Optional[int]): Maximum length of each answer.
            temperature (Optional[float]): Sampling temperature (``None`` keeps the model default).
        """
        self.provider = ai_client.provider if ai_client is not None else resolve_provider(provider)
        self.api_key = api_key
        self.model_name = model_name or (ai_client.model_name if ai_client is not None else None)
        self.max_output_tokens = max_output_tokens
        self.temperature = temperature
        self._ai_client = ai_client

        self.prompt_types = {
            "PRL": """
            Genera preguntas tipo test de opción múltiple enfocadas en Prevención de Riesgos Laborales (PRL), basadas en el siguiente texto, siguiendo estas instrucciones iniciales:
            1.  **Centrarse en leyes y reglamentos:** Las preguntas deben poder entenderse por sí mismas, haciendo referencia únicamente a leyes o reglamentos relevantes. No se puede poner: "Según el texto"
            2.  **Evita preguntas sobre números de artículos o códigos:** No incluyas preguntas que hagan referencia a números de artículos, códigos de leyes o reglamentos.
            Además, ten en cuenta las siguientes indicaciones adicionales
            """,
            "PM": """
            Genera preguntas tipo test de opción múltiple enfocadas en Gestión de Proyectos.
            """
        }
        self.existing_bank_df = None
        self.similarity_threshold = 0.8  # Default threshold
        self.current_prompt_example_content_type = "solo_enunciados"  # Default value

        self.final_df = pd.DataFrame()
        self.rejected_questions_list: List[Dict] = []
        self.invalid_questions_list: List[Dict] = []
        self.last_output_paths: Dict[str, str] = {}
        self.cancelled = False
        #: Questions discarded because they repeated another one (same statement and options).
        self.duplicate_questions_count = 0
        #: Last non-fatal provider error of the last run (shown by the GUI when nothing is generated).
        self.last_provider_error: Optional[str] = None
        self._progress_callback: Optional[Callable[[str], None]] = None
        self._bank_index_source = None
        self._bank_index: List[Tuple[set, str]] = []

    # ------------------------------------------------------------------
    # Provider access
    # ------------------------------------------------------------------
    @property
    def ai_client(self) -> BaseAIClient:
        """Client of the configured provider (created on first use)."""
        if self._ai_client is None:
            self._ai_client = create_ai_client(
                self.provider, self.api_key, self.model_name,
                max_output_tokens=self.max_output_tokens, temperature=self.temperature,
            )
        return self._ai_client

    @property
    def client(self) -> Any:
        """Underlying SDK client (kept for backward compatibility)."""
        return self.ai_client.sdk_client

    @client.setter
    def client(self, sdk_client: Any) -> None:
        # Previous versions allowed replacing the Gemini SDK client (e.g. with a mock).
        self._ai_client = create_ai_client(
            self.provider, self.api_key, self.model_name, sdk_client=sdk_client,
            max_output_tokens=self.max_output_tokens, temperature=self.temperature,
        )

    def _log(self, message: str) -> None:
        print(message)
        if self._progress_callback is not None:
            try:
                self._progress_callback(message)
            except Exception:
                pass

    # ------------------------------------------------------------------
    # Source loading
    # ------------------------------------------------------------------
    def extract_pdf_text(self, pdf_path: str) -> Optional[str]:
        """
        Extracts text from a PDF file.

        Args:
            pdf_path (str): The path to the PDF file.

        Returns:
            Optional[str]: The extracted text from the PDF, or None if an error occurs.
        """
        self.text = ""
        self.page_texts = []  # List to store the text of each page
        self.num_pages = 0
        try:
            with open(pdf_path, 'rb') as pdf_file:
                pdf_reader = PdfReader(pdf_file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text() or ""
                    self.page_texts.append(page_text)
                    self.text += page_text
                self.num_pages = len(pdf_reader.pages)  # Save the total number of pages
            return self.text
        except Exception as e:
            print(f"Error extracting text from PDF: {e}")
            return None

    def extract_topic_number_from_path(self, pdf_path: str) -> str:
        """
        Extracts the topic name from the PDF filename.

        Args:
            pdf_path (str): The path to the PDF file.

        Returns:
            str: The file name without its extension, used as the topic.
        """
        file_name = os.path.basename(pdf_path)
        self.topic_number = os.path.splitext(file_name)[0] or file_name
        return self.topic_number

    @staticmethod
    def _is_image_path(path: str) -> bool:
        """Returns True if path has a known image extension."""
        return os.path.splitext(path)[1].lower() in IMAGE_EXTENSIONS

    @staticmethod
    def _guess_mime_type(path: str) -> str:
        """Guesses MIME type from file extension for image parts."""
        ext = os.path.splitext(path)[1].lower()
        if ext in (".jpg", ".jpeg"):
            return "image/jpeg"
        if ext == ".webp":
            return "image/webp"
        if ext == ".bmp":
            return "image/bmp"
        if ext in (".tif", ".tiff"):
            return "image/tiff"
        if ext == ".gif":
            return "image/gif"
        return "image/png"

    @staticmethod
    def _group_items_into_total_chunks(items: List[Dict], total_chunks: int) -> List[List[Dict]]:
        """Splits a list into a fixed number of balanced chunks."""
        if not items or total_chunks <= 0:
            return []

        total_chunks = min(total_chunks, len(items))
        base_size = len(items) // total_chunks
        remainder = len(items) % total_chunks

        chunks = []
        start = 0
        for idx in range(total_chunks):
            extra = 1 if idx < remainder else 0
            end = start + base_size + extra
            chunks.append(items[start:end])
            start = end
        return chunks

    @staticmethod
    def _group_items_by_chunk_size(items: List[Dict], chunk_size: int) -> List[List[Dict]]:
        """Splits a list into chunks of a given size."""
        if not items or chunk_size <= 0:
            return []
        return [items[i:i + chunk_size] for i in range(0, len(items), chunk_size)]

    @staticmethod
    def _page_ranges(num_pages: int, process_by_pages: bool, pages_per_chunk: int,
                     total_chunks: Optional[int]) -> List[Tuple[int, int]]:
        """Returns the ``[start, end)`` page ranges of each chunk."""
        if num_pages <= 0:
            return []
        if not process_by_pages or num_pages == 1:
            return [(0, num_pages)]
        if total_chunks:
            pages = list(range(num_pages))
            groups = QuestionGenerator._group_items_into_total_chunks(pages, total_chunks)
        else:
            groups = [list(range(start, min(start + max(1, pages_per_chunk), num_pages)))
                      for start in range(0, num_pages, max(1, pages_per_chunk))]
        return [(group[0], group[-1] + 1) for group in groups if group]

    @staticmethod
    def _pages_label(start: int, end: int) -> str:
        return f"página {start + 1}" if end - start == 1 else f"páginas {start + 1}-{end}"

    def _count_pdf_pages(self, pdf_path: str) -> int:
        if fitz is not None:
            with fitz.open(pdf_path) as doc:
                return doc.page_count
        with open(pdf_path, "rb") as pdf_file:
            return len(PdfReader(pdf_file).pages)

    def _extract_pdf_pages_as_images(self, pdf_path: str, zoom: float = 2.0) -> List[Dict]:
        """Renders each PDF page as PNG bytes for multimodal generation."""
        if fitz is None:
            raise ImportError(
                "PyMuPDF no está instalado. Instala 'PyMuPDF' para usar input_mode='image' con PDFs."
            )

        rendered_pages: List[Dict] = []
        doc = fitz.open(pdf_path)
        matrix = fitz.Matrix(zoom, zoom)
        stem = os.path.splitext(os.path.basename(pdf_path))[0]
        try:
            for page_number, page in enumerate(doc, start=1):
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                rendered_pages.append({
                    "bytes": pix.tobytes("png"),
                    "mime_type": "image/png",
                    "name": f"{stem}_p{page_number}.png",
                })
        finally:
            doc.close()

        return rendered_pages

    def _load_source_as_image_items(self, source_path: str) -> List[Dict]:
        """Loads an image file or a PDF file (rendered pages) as image attachments."""
        if source_path.lower().endswith(".pdf"):
            return self._extract_pdf_pages_as_images(source_path)

        if not self._is_image_path(source_path):
            raise ValueError(
                f"Formato no soportado para input_mode='image': {source_path}. "
                "Use PDF o imágenes (png/jpg/webp/bmp/gif/tiff)."
            )

        extension = os.path.splitext(source_path)[1].lower()
        if extension in _IMAGE_EXTENSIONS_TO_CONVERT and fitz is not None:
            # BMP, TIFF (also multi-page) and GIF are not accepted by every provider: sent as PNG.
            stem = os.path.splitext(os.path.basename(source_path))[0]
            with fitz.open(source_path) as image_doc:
                return [{
                    "bytes": page.get_pixmap(alpha=False).tobytes("png"),
                    "mime_type": "image/png",
                    "name": f"{stem}.png" if image_doc.page_count == 1 else f"{stem}_p{number}.png",
                } for number, page in enumerate(image_doc, start=1)]

        with open(source_path, "rb") as f:
            image_bytes = f.read()
        return [{
            "bytes": image_bytes,
            "mime_type": self._guess_mime_type(source_path),
            "name": os.path.basename(source_path),
        }]

    def _pdf_fragment_bytes(self, pdf_path: str, start: int, end: int, num_pages: int) -> bytes:
        """Returns the bytes of a PDF containing only the pages ``[start, end)``."""
        if start == 0 and end >= num_pages:
            with open(pdf_path, "rb") as f:
                return f.read()
        if fitz is None:
            raise ImportError("PyMuPDF no está instalado. Instala 'PyMuPDF' para fragmentar PDFs.")
        with fitz.open(pdf_path) as source, fitz.open() as fragment:
            fragment.insert_pdf(source, from_page=start, to_page=end - 1)
            return fragment.tobytes()

    def _group_pages_into_chunks(self, pages_per_chunk: int) -> List[str]:
        """
        Groups the text from pages into larger chunks.

        Args:
            pages_per_chunk (int): The number of pages to include in each chunk.

        Returns:
            List[str]: A list of strings, where each string is a text chunk.
        """
        if not hasattr(self, 'page_texts') or not self.page_texts:
            return []

        text_chunks = []
        current_chunk_pages = []
        for i, page_text in enumerate(self.page_texts):
            current_chunk_pages.append(page_text)
            if (i + 1) % pages_per_chunk == 0 or (i + 1) == len(self.page_texts):
                text_chunks.append("\n".join(current_chunk_pages))
                current_chunk_pages = []
        return text_chunks

    def _group_pages_into_total_chunks(self, total_chunks: int) -> List[str]:
        """Groups page texts into a fixed number of balanced chunks."""
        if not hasattr(self, 'page_texts') or not self.page_texts:
            return []
        if total_chunks <= 0:
            return []

        num_pages = len(self.page_texts)
        total_chunks = min(total_chunks, num_pages)
        base_size = num_pages // total_chunks
        remainder = num_pages % total_chunks

        chunks = []
        start = 0
        for idx in range(total_chunks):
            extra = 1 if idx < remainder else 0
            end = start + base_size + extra
            chunks.append("\n".join(self.page_texts[start:end]))
            start = end
        return chunks

    @staticmethod
    def _group_text_by_length(text: str, total_chunks: int) -> List[str]:
        """Splits text into a fixed number of near-balanced chunks by word count."""
        if not text or total_chunks <= 0:
            return []
        words = text.split()
        if not words:
            return []

        total_chunks = min(total_chunks, len(words))
        base_size = len(words) // total_chunks
        remainder = len(words) % total_chunks

        chunks = []
        start = 0
        for idx in range(total_chunks):
            extra = 1 if idx < remainder else 0
            end = start + base_size + extra
            chunks.append(" ".join(words[start:end]))
            start = end
        return chunks

    @staticmethod
    def _validate_modes(input_mode: str, chunking_mode: str) -> None:
        if input_mode not in INPUT_MODES:
            raise ValueError(f"input_mode debe ser uno de: {', '.join(INPUT_MODES)}.")
        if chunking_mode not in CHUNKING_MODES:
            raise ValueError("chunking_mode debe ser 'pages' o 'text_length'.")

    def _plan_text_chunks(self, pdf_path: str, process_by_pages: bool, pages_per_chunk: int,
                          total_chunks: Optional[int], chunking_mode: str) -> Tuple[List[Dict], Optional[str]]:
        """Splits the text of a PDF into chunks. Returns ``(chunks, full_text)``."""
        pdf_full_text_content = self.extract_pdf_text(pdf_path)
        if not pdf_full_text_content or not pdf_full_text_content.strip():
            self._log(
                f"No se pudo extraer texto del PDF: {pdf_path}. Si es un PDF escaneado, "
                "usa el modo de entrada 'image' o 'pdf'."
            )
            return [], None

        if not (process_by_pages and self.page_texts):
            self._log("Procesando el PDF completo como un solo fragmento.")
            return [{"text": pdf_full_text_content, "label": "documento completo", "attachments": []}], \
                pdf_full_text_content

        if chunking_mode == "text_length":
            inferred_chunks = total_chunks
            if inferred_chunks is None:
                inferred_chunks = max(1, math.ceil(len(self.page_texts) / max(1, pages_per_chunk)))
            texts = self._group_text_by_length(pdf_full_text_content, inferred_chunks)
            chunks = [{"text": text, "label": f"parte {i + 1} de {len(texts)}", "attachments": []}
                      for i, text in enumerate(texts)]
            self._log(f"PDF dividido en {len(chunks)} fragmentos (modo: {chunking_mode}).")
        else:
            ranges = self._page_ranges(len(self.page_texts), True, pages_per_chunk, total_chunks)
            chunks = [{"text": "\n".join(self.page_texts[start:end]), "label": self._pages_label(start, end),
                       "attachments": []} for start, end in ranges]
            empty = [chunk["label"] for chunk in chunks if not chunk["text"].strip()]
            if empty:
                self._log(f"Se omiten fragmentos sin texto extraíble ({', '.join(empty)}); para páginas escaneadas "
                          "usa el modo 'image' o 'pdf'.")
                chunks = [chunk for chunk in chunks if chunk["text"].strip()]
            if total_chunks is None:
                self._log(f"PDF dividido en {len(chunks)} fragmentos de ~{pages_per_chunk} página(s).")
            else:
                self._log(f"PDF dividido en {len(chunks)} fragmentos (modo: {chunking_mode}).")
        return chunks, pdf_full_text_content

    def _plan_attachment_chunks(self, source_path: str, input_mode: str, process_by_pages: bool,
                                pages_per_chunk: int, total_chunks: Optional[int], load_bytes: bool,
                                attachments_dir: Optional[str] = None,
                                fragment_stem: Optional[str] = None) -> List[Dict]:
        """
        Splits an image/PDF source into chunks of attachments for the 'image' and 'pdf' modes.

        Each chunk is ``{"label", "attachments", "files"}``. ``attachments`` holds the bytes to
        send through the API (only when ``load_bytes``); ``files`` lists the files the user must
        attach in a web chat (fragments are written to ``attachments_dir`` when given, named after
        ``fragment_stem`` or, by default, after the source file).
        """
        name = os.path.basename(source_path)
        if not source_path.lower().endswith(".pdf"):
            if not self._is_image_path(source_path):
                raise ValueError(
                    f"Formato no soportado para input_mode='{input_mode}': {source_path}. "
                    "Use PDF o imágenes (png/jpg/webp/bmp/gif/tiff)."
                )
            attachments = self._load_source_as_image_items(source_path) if load_bytes else []
            return [{"label": "imagen", "attachments": attachments, "files": [source_path]}]

        needs_pages = process_by_pages or (load_bytes and input_mode == "image")
        if not needs_pages and not attachments_dir:
            # Manual workflow with the whole document: the file does not need to be opened.
            attachments = []
            if load_bytes:
                with open(source_path, "rb") as f:
                    attachments = [{"bytes": f.read(), "mime_type": "application/pdf", "name": name}]
            return [{"label": "documento completo", "attachments": attachments, "files": [source_path]}]

        num_pages = self._count_pdf_pages(source_path)
        ranges = self._page_ranges(num_pages, process_by_pages, pages_per_chunk, total_chunks)
        page_images = self._extract_pdf_pages_as_images(source_path) if load_bytes and input_mode == "image" else []
        stem = fragment_stem or os.path.splitext(name)[0]

        chunks = []
        for start, end in ranges:
            whole_document = start == 0 and end >= num_pages
            pages_suffix = f"p{start + 1}" if end - start == 1 else f"p{start + 1}-{end}"
            fragment_name = name if whole_document else f"{stem}_{pages_suffix}.pdf"
            attachments: List[Dict] = []
            if load_bytes:
                if input_mode == "image":
                    attachments = page_images[start:end]
                else:
                    attachments = [{
                        "bytes": self._pdf_fragment_bytes(source_path, start, end, num_pages),
                        "mime_type": "application/pdf",
                        "name": fragment_name,
                    }]
            files = [source_path]
            if attachments_dir and not whole_document:
                os.makedirs(attachments_dir, exist_ok=True)
                fragment_path = os.path.join(attachments_dir, fragment_name)
                with open(fragment_path, "wb") as f:
                    f.write(self._pdf_fragment_bytes(source_path, start, end, num_pages))
                files = [fragment_path]
            chunks.append({
                "label": "documento completo" if whole_document else self._pages_label(start, end),
                "attachments": attachments,
                "files": files,
            })
        if len(chunks) > 1:
            self._log(f"Fuente dividida en {len(chunks)} fragmentos.")
        return chunks

    @staticmethod
    def _attachment_instruction(input_mode: str, file_names: List[str], label: str) -> str:
        names = ", ".join(file_names)
        scope = "" if label in ("documento completo", "imagen") else f" ({label} del documento original)"
        if input_mode == "image" and any(name.lower().endswith(".pdf") for name in file_names):
            # Manual workflow: the user attaches the PDF (or fragment) itself to the chat.
            return (
                f"IMPORTANTE: Debes basarte ÚNICAMENTE en las páginas del documento adjunto ({names}){scope}, "
                "analizando su contenido visual (texto, tablas, esquemas e imágenes). "
                "No inventes información que no aparezca en esas páginas."
            )
        if input_mode == "image":
            return (
                "IMPORTANTE: Debes basarte ÚNICAMENTE en las imágenes adjuntas"
                f"{f' ({names})' if names else ''}{scope}. "
                "No inventes información que no sea visible en esas imágenes."
            )
        return (
            "IMPORTANTE: Debes basarte ÚNICAMENTE en el contenido del documento adjunto"
            f"{f' ({names})' if names else ''}{scope}. "
            "No inventes información que no aparezca en ese documento."
        )

    # ------------------------------------------------------------------
    # Prompt construction
    # ------------------------------------------------------------------
    def _build_prompt(
            self,
            working_chunk_text: str,
            num_questions_to_generate: int,
            prompt_type: str,
            custom_prompt: Optional[str],
            existing_questions_df: Optional[pd.DataFrame],
            current_topic: str,
            full_document_context_text: Optional[str] = None,
            bank_prompt_scope: Optional[str] = None,
            questions_from_current_chunk_attempts: Optional[List[Dict]] = None,
            for_web_chat: bool = False,
            material_is_attachment: bool = False,
    ) -> str:
        """
        Constructs the complete prompt to be sent to the model.

        Args:
            working_chunk_text (str): The specific text segment to generate questions from
                (or the instruction describing the attached material).
            num_questions_to_generate (int): The target number of questions to generate.
            prompt_type (str): The key for the predefined prompt type (e.g., "PRL").
            custom_prompt (Optional[str]): A user-provided custom prompt.
            existing_questions_df (Optional[pd.DataFrame]): DataFrame of existing questions for guidance.
            current_topic (str): The topic of the current PDF.
            full_document_context_text (Optional[str]): The full text of the document for context.
            bank_prompt_scope (Optional[str]): Scope for selecting guidance questions ('mismo_tema', 'todos_los_temas').
            questions_from_current_chunk_attempts (Optional[List[Dict]]): Questions already generated for this chunk.
            for_web_chat (bool): If True, asks for the JSON inside a Markdown ``json`` code block (easy to copy
                from Google AI Studio, ChatGPT or Claude).
            material_is_attachment (bool): If True, ``working_chunk_text`` describes attached files.

        Returns:
            str: The fully constructed prompt string.
        """
        if custom_prompt:
            user_prefix = custom_prompt
        elif prompt_type in self.prompt_types:
            user_prefix = self.prompt_types[prompt_type]
        else:
            user_prefix = "Genera preguntas tipo test de opción múltiple basadas en el siguiente texto, siguiendo estas instrucciones:"
        user_prefix = textwrap.dedent(user_prefix).strip()

        fixed_prefix = (
            f"Genera {num_questions_to_generate} preguntas tipo test de opción múltiple con una única "
            "respuesta correcta basadas en el material que adjunto al final, siguiendo estas instrucciones:"
        )

        # --- Guidance from the existing question bank ---
        prompt_guidance_from_bank_str = ""
        if existing_questions_df is not None and \
                not existing_questions_df.empty and \
                bank_prompt_scope and bank_prompt_scope != "no_incluir_en_prompt":

            questions_for_guidance_df = pd.DataFrame()
            header_for_bank_guidance = ""

            if bank_prompt_scope == "mismo_tema":
                if 'Tema' in existing_questions_df.columns:
                    questions_for_guidance_df = existing_questions_df[
                        existing_questions_df['Tema'].astype(str).str.strip().str.lower() == current_topic.strip().lower()
                        ]
                if not questions_for_guidance_df.empty:
                    header_for_bank_guidance = "Para evitar redundancia con el BANCO DE PREGUNTAS EXISTENTE, considera las siguientes preguntas del MISMO TEMA (asegúrate de que las nuevas preguntas sean originales y basadas estrictamente en el material proporcionado):\n"
            elif bank_prompt_scope == "todos_los_temas":
                questions_for_guidance_df = existing_questions_df
                if not questions_for_guidance_df.empty:
                    header_for_bank_guidance = "Para evitar redundancia con el BANCO DE PREGUNTAS EXISTENTE, considera las siguientes preguntas de TODO EL BANCO (asegúrate de que las nuevas preguntas sean originales y basadas estrictamente en el material proporcionado):\n"

            if not questions_for_guidance_df.empty:
                prompt_guidance_from_bank_str += header_for_bank_guidance
                for _, row in questions_for_guidance_df.iterrows():
                    question_text = row.get('Pregunta', '')
                    if not isinstance(question_text, str) or not question_text.strip():
                        continue
                    prompt_guidance_from_bank_str += f"- Pregunta del banco: {question_text.strip()}\n"
                    if self.current_prompt_example_content_type == "enunciados_y_respuestas":
                        options = [row.get(f'Respuesta {letter}', '') for letter in OPTION_LETTERS]
                        options = [str(opt).strip() for opt in options if opt is not None and pd.notna(opt)]
                        options_str = "\n".join(f"  - {opt}" for opt in options if opt)
                        if options_str:
                            prompt_guidance_from_bank_str += f"{options_str}\n"

        # --- Guidance from questions already generated for this chunk in previous attempts ---
        prompt_guidance_from_current_attempts_str = ""
        if questions_from_current_chunk_attempts:  # This is a list of dictionaries
            prompt_guidance_from_current_attempts_str += "Adicionalmente, para evitar generar preguntas idénticas o muy similares a las YA GENERADAS PARA ESTE MISMO FRAGMENTO en intentos previos, ten en cuenta las siguientes preguntas que ya se han aceptado:\n"
            for question_dict in questions_from_current_chunk_attempts:
                question_text = question_dict.get('Pregunta', '')
                if question_text:
                    prompt_guidance_from_current_attempts_str += f"- Pregunta ya generada para este fragmento: {question_text}\n"

        if full_document_context_text is not None:
            context_section = (
                "Aquí está el **CONTEXTO COMPLETO** del documento, que te proporcionamos para que tengas una "
                "comprensión general del tema y el documento en su totalidad. Puedes referenciar este texto para "
                "obtener contexto adicional, pero las preguntas deben basarse estrictamente en el **FRAGMENTO DE TRABAJO**.\n"
                f"{full_document_context_text}"
            )
            chunk_instruction_and_text_section = (
                "---\n"
                "**IMPORTANTE**: Debes generar las preguntas **ÚNICAMENTE** basándote en el siguiente "
                "\"**FRAGMENTO DE TRABAJO**\". La respuesta a cada pregunta debe ser directamente derivable del "
                "contenido de este FRAGMENTO DE TRABAJO.\n"
                "---\n"
                "Aquí está el FRAGMENTO DE TRABAJO sobre el que debes generar las preguntas:\n"
                f"{working_chunk_text}"
            )
        elif material_is_attachment:
            context_section = ""
            chunk_instruction_and_text_section = f"**MATERIAL**: {working_chunk_text}"
        else:
            context_section = ""
            chunk_instruction_and_text_section = (
                "Aquí está el **TEXTO COMPLETO** sobre el que debes generar las preguntas:\n"
                f"{working_chunk_text}"
            )

        sections = [
            user_prefix,
            fixed_prefix,
            _FORMAT_INSTRUCTIONS,
            _OUTPUT_RULE_WEB_CHAT if for_web_chat else _OUTPUT_RULE_API,
            prompt_guidance_from_bank_str.strip(),
            prompt_guidance_from_current_attempts_str.strip(),
            context_section,
            chunk_instruction_and_text_section,
            # After a long material the model may forget the output rule given at the beginning.
            _FINAL_REMINDER_WEB_CHAT if for_web_chat else _FINAL_REMINDER_API,
        ]
        return "\n\n".join(section for section in sections if section)

    # ------------------------------------------------------------------
    # Response analysis
    # ------------------------------------------------------------------
    def _analyze_response(self, response_text: str, topic_number: str) -> List[Dict]:
        """
        Parses the model answer, filters it by similarity and returns the accepted questions.

        Questions similar to the existing bank are added to ``self.rejected_questions_list``
        and fragments that cannot be used (e.g. 3 options, no valid answer letter) to
        ``self.invalid_questions_list``.

        Args:
            response_text (str): The raw text answer (API or pasted from a web chat).
            topic_number (str): The topic associated with these questions.

        Returns:
            List[Dict]: A list of accepted question dictionaries (``QUESTION_COLUMNS`` keys).
        """
        questions, errors = parse_questions_response(response_text or "")
        for error in errors:
            self.invalid_questions_list.append({"Tema": f"{topic_number}", **error})

        accepted_questions = []
        for question in questions:
            question_dict = {
                "Tema": f"{topic_number}",
                "Estado": "Pendiente de revisar",
                "Pregunta": question["Pregunta"],
                **{f"Respuesta {letter}": option for letter, option in zip(OPTION_LETTERS, question["Opciones"])},
                "Respuesta correcta": question["Respuesta correcta"],
                "Texto relevante": question["Texto relevante"],
            }
            is_similar, sim_score, matched_text = self._is_similar_question(question["Pregunta"])
            if is_similar:
                question_dict['Similitud'] = sim_score
                question_dict['Pregunta Coincidente'] = matched_text
                self.rejected_questions_list.append(question_dict)
            else:
                accepted_questions.append(question_dict)

        if not questions and not errors and (response_text or "").strip():
            self.invalid_questions_list.append({
                "Tema": f"{topic_number}",
                "Fragmento": response_text.strip()[:300],
                "Motivo": "No se encontraron preguntas en formato JSON en la respuesta.",
            })
        return accepted_questions

    def _analyze_gemini_response(self, response_text: str, topic_number: str) -> List[Dict]:
        """Backward-compatible alias of :meth:`_analyze_response`."""
        return self._analyze_response(response_text, topic_number)

    def parse_questions_from_text(self, response_text: str, topic: str = "Sin tema",
                                  apply_similarity_filter: bool = True) -> pd.DataFrame:
        """
        Converts a model answer (e.g. copied from Google AI Studio, ChatGPT or Claude) into a
        DataFrame with the question-bank columns, without saving anything.

        Args:
            response_text (str): Answer text (JSON array, Markdown ``json`` code block, legacy format...).
            topic (str): Value for the 'Tema' column.
            apply_similarity_filter (bool): If True, questions similar to ``existing_bank_df``
                are moved to ``rejected_questions_list``.

        Returns:
            pd.DataFrame: The valid questions.
        """
        previous_bank = self.existing_bank_df
        if not apply_similarity_filter:
            self.existing_bank_df = None
        try:
            rows = self._analyze_response(response_text, topic)
        finally:
            self.existing_bank_df = previous_bank
        return pd.DataFrame(rows, columns=QUESTION_COLUMNS)

    # ------------------------------------------------------------------
    # Output files
    # ------------------------------------------------------------------
    def save_questions_to_excel(
            self,
            df: pd.DataFrame,
            filepath: str
    ) -> None:
        """Saves the questions to an Excel file."""
        if not df.empty:
            if 'Número de pregunta' not in df.columns:
                df.insert(0, 'Número de pregunta', range(1, len(df) + 1))
            df.to_excel(filepath, index=False)
            print(f"Preguntas aceptadas guardadas en: '{filepath}'")
        else:
            print("No se generaron preguntas aceptadas para guardar en Excel.")

    def save_questions_to_docx(
            self,
            df: pd.DataFrame,
            filepath: str
    ) -> None:
        """Saves the questions to a DOCX file for manual review."""
        if not df.empty:
            document = Document()
            document.add_heading('Banco de Preguntas - Pendiente de Revisar', level=1)
            for number, (_, row) in enumerate(df.iterrows(), start=1):
                pregunta_paragraph = document.add_paragraph()
                pregunta_run = pregunta_paragraph.add_run(f"Pregunta {number}: {row['Pregunta']}")
                pregunta_run.bold = True
                document.add_paragraph(f"Tema: {row['Tema']}")
                document.add_paragraph(f"Estado: {row['Estado']}")
                document.add_paragraph(f"A) {row['Respuesta A']}")
                document.add_paragraph(f"B) {row['Respuesta B']}")
                document.add_paragraph(f"C) {row['Respuesta C']}")
                document.add_paragraph(f"D) {row['Respuesta D']}")
                document.add_paragraph(f"Respuesta correcta: {row['Respuesta correcta']}")
                document.add_paragraph(f"Texto relevante: {row['Texto relevante']}")
                document.add_paragraph()
            document.save(filepath)
        else:
            print("No hay preguntas para guardar en DOCX.")

    def save_rejected_questions_to_docx(self, df: pd.DataFrame, filepath: str) -> None:
        """
        Saves questions discarded due to similarity to a DOCX file for inspection.

        Args:
            df (pd.DataFrame): The DataFrame containing the discarded questions and similarity details.
            filepath (str): The filepath for the output file.
        """
        if df.empty:
            print("No se descartaron preguntas por similitud.")
            return

        document = Document()
        document.add_heading('Preguntas Descartadas por Similitud', level=1)
        document.add_paragraph(
            "Las siguientes preguntas fueron generadas por el modelo, pero se descartaron "
            "automáticamente porque su similitud con una pregunta existente en el banco "
            "superó el umbral establecido."
        )

        for index, row in df.iterrows():
            # Visual separator between questions
            document.add_paragraph("---")

            # Information about the discarded question
            p = document.add_paragraph()
            p.add_run('Pregunta descartada: ').bold = True
            p.add_run(f"{row['Pregunta']}")

            document.add_paragraph(f"Tema: {row['Tema']}")
            document.add_paragraph(f"A) {row['Respuesta A']}")
            document.add_paragraph(f"B) {row['Respuesta B']}")
            document.add_paragraph(f"C) {row['Respuesta C']}")
            document.add_paragraph(f"D) {row['Respuesta D']}")
            document.add_paragraph(f"Respuesta correcta: {row['Respuesta correcta']}")

            # Information about the reason for rejection
            p_reason = document.add_paragraph()
            p_reason.add_run('Motivo del descarte:').bold = True

            # Format the similarity score as a percentage
            similarity_score_pct = f"{row.get('Similitud', 0.0) * 100:.1f}%"

            p_details = document.add_paragraph()
            p_details.add_run(f"Similitud: ").bold = True
            p_details.add_run(f"{similarity_score_pct}")

            p_match = document.add_paragraph()
            p_match.add_run("Coincide con la pregunta existente: ").bold = True
            p_match.add_run(f"{row.get('Pregunta Coincidente', 'N/A')}")

        document.save(filepath)
        print(f"Se guardó un informe de {len(df)} preguntas descartadas en: '{filepath}'")

    def _save_generation_outputs(self, output_filename: str, output_dir: str, generate_docx: bool) -> Dict[str, str]:
        """Writes the accepted, rejected and invalid questions. Returns the written paths."""
        os.makedirs(output_dir, exist_ok=True)
        paths: Dict[str, str] = {}
        excel_output_path = os.path.join(output_dir, f'{output_filename}_pendiente_de_revisar.xlsx')
        docx_output_path = os.path.join(output_dir, f'{output_filename}_pendiente_de_revisar.docx')

        self.save_questions_to_excel(self.final_df, excel_output_path)
        if not self.final_df.empty:
            paths["xlsx"] = excel_output_path
        if generate_docx and not self.final_df.empty:
            self.save_questions_to_docx(self.final_df, docx_output_path)
            paths["docx"] = docx_output_path

        if self.rejected_questions_list:
            rejected_df = pd.DataFrame(self.rejected_questions_list)
            rejected_docx_path = os.path.join(output_dir, f'{output_filename}_similares_descartadas.docx')
            self.save_rejected_questions_to_docx(rejected_df, rejected_docx_path)
            paths["rejected_docx"] = rejected_docx_path
        else:
            print("No se descartaron preguntas por similitud durante todo el proceso.")

        if self.invalid_questions_list:
            invalid_path = os.path.join(output_dir, f'{output_filename}_no_validas.xlsx')
            pd.DataFrame(self.invalid_questions_list).to_excel(invalid_path, index=False)
            paths["invalid_xlsx"] = invalid_path
            print(f"Se guardaron {len(self.invalid_questions_list)} fragmentos no válidos en: '{invalid_path}'")

        self.last_output_paths = paths
        return paths

    # ------------------------------------------------------------------
    # Similarity filter
    # ------------------------------------------------------------------
    def _normalize_text(self, text: str) -> str:
        """Normalizes text for comparison (lowercase, no accents, no punctuation)."""
        if not isinstance(text, str):
            return ""
        text = unicodedata.normalize("NFKD", text.lower())
        text = "".join(ch for ch in text if not unicodedata.combining(ch))
        text = re.sub(r'[^\w\s]', '', text)  # Remove punctuation
        return text

    def _calculate_jaccard_similarity(self, text1: str, text2: str) -> float:
        """Calculates the Jaccard similarity between two texts."""
        words1 = set(self._normalize_text(text1).split())
        words2 = set(self._normalize_text(text2).split())
        return self._jaccard(words1, words2)

    @staticmethod
    def _jaccard(words1: set, words2: set) -> float:
        if not words1 and not words2:
            return 1.0
        if not words1 or not words2:
            return 0.0
        return len(words1 & words2) / len(words1 | words2)

    def _load_existing_bank(self, existing_bank_path: Optional[str]) -> None:
        """Loads the bank used for the similarity filter and the prompt guidance."""
        self.existing_bank_df = None
        if not existing_bank_path:
            return
        if not os.path.exists(existing_bank_path):
            self._log(f"Advertencia: El archivo del banco de preguntas existente no se encontró en '{existing_bank_path}'.")
            return
        try:
            self.existing_bank_df = pd.read_excel(existing_bank_path)
            self._log(f"Banco de preguntas existente cargado desde '{existing_bank_path}' para el filtro de similitud.")
        except Exception as e:
            self._log(f"Error al cargar el banco de preguntas existente: {e}")
            self.existing_bank_df = None

    def _get_bank_index(self) -> List[Tuple[set, str]]:
        """Word sets of the bank statements (cached while ``existing_bank_df`` does not change)."""
        if self._bank_index_source is not self.existing_bank_df:
            index: List[Tuple[set, str]] = []
            if self.existing_bank_df is not None and 'Pregunta' in self.existing_bank_df.columns:
                for text in self.existing_bank_df['Pregunta']:
                    if isinstance(text, str) and text.strip():
                        index.append((set(self._normalize_text(text).split()), text))
            self._bank_index = index
            self._bank_index_source = self.existing_bank_df
        return self._bank_index

    def _is_similar_question(self,
                             new_question_text: str
                             ) -> Tuple[bool, float, Optional[str]]:
        """
        Checks if the new question is similar to any existing one in the bank,
        comparing ONLY the question statements for filtering purposes.

        Args:
            new_question_text (str): The text of the new question.

        Returns:
            Tuple[bool, float, Optional[str]]: (is_similar, similarity_score, matching_question_text)
        """
        if self.existing_bank_df is None or self.existing_bank_df.empty or self.similarity_threshold > 1.0:
            # If the threshold is > 1.0 (e.g., 1.1), the filter is effectively disabled
            return (False, 0.0, None)

        new_words = set(self._normalize_text(new_question_text).split())
        for existing_words, existing_text in self._get_bank_index():
            question_similarity = self._jaccard(new_words, existing_words)
            if question_similarity >= self.similarity_threshold:
                # For filtering, the similarity of the statement is sufficient
                return (True, question_similarity, existing_text)

        return (False, 0.0, None)

    def _set_similarity_threshold(self, similarity_threshold: Optional[float]) -> None:
        if similarity_threshold is not None and 0 <= similarity_threshold <= 1.0:
            self.similarity_threshold = similarity_threshold
            self._log(f"Filtro de similitud activado con umbral: {self.similarity_threshold}")
        else:
            self.similarity_threshold = 1.1  # Effectively disables the filter
            self._log(f"Filtro de similitud DESACTIVADO (umbral: {self.similarity_threshold}).")

    def filter_generated_questions_duplicates(
            self,
            generated_questions_df: pd.DataFrame,
            existing_bank_path: str,
            similarity_threshold: float = 0.8
    ) -> Tuple[pd.DataFrame, pd.DataFrame]:
        """
        Filters an external DataFrame of generated questions against an existing bank.

        Returns:
            Tuple[pd.DataFrame, pd.DataFrame]: (accepted_df, rejected_df)
        """
        if generated_questions_df is None or generated_questions_df.empty:
            return pd.DataFrame(), pd.DataFrame()

        if not existing_bank_path or not os.path.exists(existing_bank_path):
            raise FileNotFoundError(f"No se encontró el banco de preguntas: {existing_bank_path}")

        self.existing_bank_df = pd.read_excel(existing_bank_path)
        self.similarity_threshold = similarity_threshold if 0 <= similarity_threshold <= 1.0 else 0.8

        accepted_rows = []
        rejected_rows = []
        for _, row in generated_questions_df.iterrows():
            question_text = row.get("Pregunta", "")
            is_similar, sim_score, matched_text = self._is_similar_question(question_text)
            if is_similar:
                row_dict = row.to_dict()
                row_dict["Similitud"] = sim_score
                row_dict["Pregunta Coincidente"] = matched_text
                rejected_rows.append(row_dict)
            else:
                accepted_rows.append(row.to_dict())

        return pd.DataFrame(accepted_rows), pd.DataFrame(rejected_rows)

    # ------------------------------------------------------------------
    # Manual workflow (no API): prompts for the web chats and import of the answers
    # ------------------------------------------------------------------
    def _validate_sources(self, pdf_paths: List[str], input_mode: str, check_exists: bool = True) -> None:
        """
        Checks the source files before any work (and before any paid API call).

        Raises:
            FileNotFoundError: Some file does not exist (only when ``check_exists``).
            ValueError: Some file has a format that ``input_mode`` cannot use.
        """
        if check_exists:
            missing = [path for path in pdf_paths if not os.path.isfile(path)]
            if missing:
                raise FileNotFoundError("No se encuentran estos archivos fuente:\n" + "\n".join(missing))
        if input_mode == "text":
            unsupported = [path for path in pdf_paths if not path.lower().endswith(".pdf")]
            hint = "En modo 'text' solo se admiten PDF; usa el modo 'image' o 'pdf' para imágenes."
        else:
            unsupported = [path for path in pdf_paths
                           if not path.lower().endswith(".pdf") and not self._is_image_path(path)]
            hint = f"En modo '{input_mode}' se admiten PDF e imágenes (png/jpg/webp/bmp/gif/tiff)."
        if unsupported:
            raise ValueError(f"Formato no soportado: {', '.join(unsupported)}. {hint}")

    def build_generation_prompts(
            self,
            pdf_paths: List[str],
            prompt_type: str,
            num_questions_per_chunk_target: int = 5,
            custom_prompt: Optional[str] = None,
            existing_bank_path: Optional[str] = None,
            process_by_pages: bool = False,
            pages_per_chunk: int = 1,
            total_chunks: Optional[int] = None,
            chunking_mode: str = "pages",
            bank_prompt_scope: Optional[str] = None,
            prompt_example_content_type: str = "solo_enunciados",
            include_similar_questions_in_prompt: bool = True,
            input_mode: str = "text",
            for_web_chat: bool = True,
            attachments_dir: Optional[str] = None,
            include_full_document_context: bool = True,
    ) -> pd.DataFrame:
        """
        Builds the prompts to paste in a web chat (Google AI Studio, ChatGPT, Claude...)
        without calling any API, using the same chunking options as the API generation.

        The arguments shared with :meth:`generate_multiple_choice_questions` (``pdf_paths``,
        ``prompt_type``, ``num_questions_per_chunk_target``, ``custom_prompt``,
        ``existing_bank_path``, ``process_by_pages``, ``pages_per_chunk``, ``total_chunks``,
        ``chunking_mode``, ``bank_prompt_scope``, ``prompt_example_content_type`` and
        ``include_similar_questions_in_prompt``) have the same meaning there.

        Args:
            input_mode (str): 'text' embeds the PDF text in the prompt; 'image' and 'pdf'
                ask the user to attach the file (or fragment) to the chat message.
            for_web_chat (bool): Ask for the answer inside a Markdown ``json`` code block (easy to copy).
            attachments_dir (Optional[str]): In 'image'/'pdf' modes with fragments, folder
                where the PDF fragment of each prompt is written so it can be attached.
            include_full_document_context (bool): In 'text' mode with fragments, also include
                the whole document as context (longer prompts).

        Returns:
            pd.DataFrame: One row per prompt with the columns 'PDF', 'Tema', 'Fragmento',
            'Total fragmentos', 'Páginas', 'Adjuntar', 'Caracteres' and 'Prompt'.
        """
        self._validate_modes(input_mode, chunking_mode)
        # Existence is checked when a file is read (in 'image'/'pdf' modes without fragments the
        # prompt only names the file).
        self._validate_sources(pdf_paths, input_mode, check_exists=False)
        prompts_rows: List[Dict] = []
        self.current_prompt_example_content_type = prompt_example_content_type
        self._load_existing_bank(existing_bank_path)
        guidance_df = self.existing_bank_df if include_similar_questions_in_prompt else None
        guidance_scope = bank_prompt_scope if include_similar_questions_in_prompt else None
        used_fragment_stems: Dict[str, int] = {}

        for pdf_path in pdf_paths:
            current_pdf_topic = self.extract_topic_number_from_path(pdf_path)

            if input_mode == "text":
                chunks, full_text = self._plan_text_chunks(
                    pdf_path, process_by_pages, pages_per_chunk, total_chunks, chunking_mode)
                use_context = process_by_pages and include_full_document_context and len(chunks) > 1
            else:
                # Two sources with the same name (T1/apuntes.pdf, T2/apuntes.pdf) must not share
                # fragment files: the second one gets the suffix _2, _3...
                stem = os.path.splitext(os.path.basename(pdf_path))[0]
                used_fragment_stems[stem] = used_fragment_stems.get(stem, 0) + 1
                fragment_stem = stem if used_fragment_stems[stem] == 1 else f"{stem}_{used_fragment_stems[stem]}"
                chunks = self._plan_attachment_chunks(
                    pdf_path, input_mode, process_by_pages, pages_per_chunk, total_chunks,
                    load_bytes=False, attachments_dir=attachments_dir, fragment_stem=fragment_stem)
                full_text, use_context = None, False

            for i, chunk in enumerate(chunks):
                files = chunk.get("files", [])
                if input_mode == "text":
                    material = chunk["text"]
                else:
                    material = self._attachment_instruction(
                        input_mode, [os.path.basename(path) for path in files], chunk["label"])
                prompt = self._build_prompt(
                    working_chunk_text=material,
                    num_questions_to_generate=num_questions_per_chunk_target,
                    prompt_type=prompt_type,
                    custom_prompt=custom_prompt,
                    existing_questions_df=guidance_df,
                    current_topic=current_pdf_topic,
                    full_document_context_text=full_text if use_context else None,
                    bank_prompt_scope=guidance_scope,
                    questions_from_current_chunk_attempts=None,
                    for_web_chat=for_web_chat,
                    material_is_attachment=input_mode != "text",
                ).strip()
                prompts_rows.append({
                    "PDF": pdf_path,
                    "Tema": current_pdf_topic,
                    "Fragmento": i + 1,
                    "Total fragmentos": len(chunks),
                    "Páginas": chunk["label"],
                    "Adjuntar": "; ".join(files) if input_mode != "text" else "",
                    "Caracteres": len(prompt),
                    "Prompt": prompt,
                })

        return pd.DataFrame(prompts_rows, columns=[
            "PDF", "Tema", "Fragmento", "Total fragmentos", "Páginas", "Adjuntar", "Caracteres", "Prompt",
        ])

    @staticmethod
    def _safe_filename(text: str, max_length: int = 60) -> str:
        cleaned = re.sub(r'[\\/:*?"<>|\s]+', "_", str(text)).strip("._")
        return cleaned[:max_length] or "prompt"

    def export_prompts(self, prompts_df: pd.DataFrame, output_dir: Optional[str] = None,
                       output_filename: Optional[str] = None,
                       responses: Optional[List[Optional[str]]] = None) -> Dict[str, str]:
        """
        Saves the prompts of :meth:`build_generation_prompts` for the manual workflow.

        Writes ``<nombre>_prompts.xlsx`` (one row per prompt with a 'Respuesta' column where the
        chat answer can be pasted) and a ``<nombre>_prompts`` folder with one ``.txt`` file per
        prompt (Excel cells are limited to ~32 000 characters).

        Args:
            prompts_df (pd.DataFrame): Result of :meth:`build_generation_prompts`.
            output_dir (Optional[str]): Output folder (current folder if None).
            output_filename (Optional[str]): Base name of the files.
            responses (Optional[List[Optional[str]]]): Answers already obtained, aligned with the
                rows of ``prompts_df`` (saved in the 'Respuesta' column, or in a ``.txt`` file
                referenced there when too long), so the work can be resumed later.

        Returns:
            Dict[str, str]: ``{"xlsx": ..., "folder": ...}``.
        """
        output_dir = output_dir or os.getcwd()
        base_name = output_filename or "preguntas"
        folder = os.path.join(output_dir, f"{base_name}_prompts")
        os.makedirs(folder, exist_ok=True)

        rows = []
        for number, (_, row) in enumerate(prompts_df.iterrows(), start=1):
            prompt = str(row.get("Prompt", ""))
            file_name = f"{number:02d}_{self._safe_filename(row.get('Tema', ''))}_frag{row.get('Fragmento', 1)}.txt"
            prompt_path = os.path.join(folder, file_name)
            with open(prompt_path, "w", encoding="utf-8") as f:
                f.write(prompt)
            response = (responses[number - 1] if responses and number - 1 < len(responses) else None) or ""
            if len(response) > _EXCEL_CELL_LIMIT:
                response_path = os.path.join(folder, file_name.replace(".txt", "_respuesta.txt"))
                with open(response_path, "w", encoding="utf-8") as f:
                    f.write(response)
                # Relative to the workbook, so the pair can be moved to another folder.
                response = os.path.relpath(response_path, output_dir)
            rows.append({
                "Nº": number,
                "Tema": row.get("Tema", ""),
                "Fragmento": row.get("Fragmento", 1),
                "Total fragmentos": row.get("Total fragmentos", 1),
                "Páginas": row.get("Páginas", ""),
                "Adjuntar": row.get("Adjuntar", ""),
                "Caracteres": len(prompt),
                "Archivo prompt": prompt_path,
                "Prompt": prompt if len(prompt) <= _EXCEL_CELL_LIMIT else
                "(Prompt demasiado largo para una celda de Excel: copia el contenido de 'Archivo prompt')",
                "Respuesta": response,
            })

        instructions = pd.DataFrame({"Instrucciones": [
            "1. Abre el chat (Google AI Studio, ChatGPT o Claude) e inicia una conversación nueva para cada prompt.",
            "2. Si la columna 'Adjuntar' indica archivos, adjúntalos al mensaje.",
            "3. Copia el prompt (columna 'Prompt' o el archivo de 'Archivo prompt') y envíalo.",
            "4. Copia la respuesta completa (botón 'Copiar' del chat) y pégala en la columna 'Respuesta'.",
            "   Si la respuesta es muy larga, pega la respuesta en un .txt y escribe su ruta en 'Respuesta'.",
            "5. En pyexamgenerator usa 'Importar respuestas' con este archivo para crear el Excel/DOCX de revisión.",
        ]})
        xlsx_path = os.path.join(output_dir, f"{base_name}_prompts.xlsx")
        with pd.ExcelWriter(xlsx_path) as writer:
            pd.DataFrame(rows).to_excel(writer, sheet_name="Prompts", index=False)
            instructions.to_excel(writer, sheet_name="Instrucciones", index=False)
        self._log(f"Prompts guardados en '{xlsx_path}' y en la carpeta '{folder}'.")
        return {"xlsx": xlsx_path, "folder": folder}

    @staticmethod
    def _normalize_responses(responses: Any, default_topic: str) -> List[Tuple[str, str]]:
        """Accepts DataFrame/dict/list/str answers and returns ``[(topic, text)]``."""
        if responses is None:
            return []
        if isinstance(responses, str):
            return [(default_topic, responses)]
        if isinstance(responses, pd.DataFrame):
            responses = responses.to_dict("records")
        if isinstance(responses, dict):
            if any(key in responses for key in ("Respuesta", "respuesta", "response")):
                responses = [responses]
            else:
                responses = list(responses.items())

        normalized = []
        for item in responses:
            if isinstance(item, str):
                topic, text = default_topic, item
            elif isinstance(item, dict):
                topic = item.get("Tema", item.get("tema", item.get("topic", default_topic)))
                text = item.get("Respuesta", item.get("respuesta", item.get("response", "")))
            else:
                topic, text = item
            if not isinstance(text, str) or not text.strip():
                continue
            if not isinstance(topic, str) or not topic.strip():
                topic = default_topic
            normalized.append((topic.strip(), text))
        return normalized

    def import_questions_from_responses(
            self,
            responses: Any,
            output_filename: Optional[str] = None,
            output_dir: Optional[str] = None,
            generate_docx: bool = True,
            existing_bank_path: Optional[str] = None,
            similarity_threshold: Optional[float] = 0.8,
            default_topic: str = "Sin tema",
    ) -> pd.DataFrame:
        """
        Imports the answers copied from a web chat (Google AI Studio, ChatGPT, Claude...) and
        saves them exactly like :meth:`generate_multiple_choice_questions` (``*_pendiente_de_revisar``
        xlsx/docx, ``*_similares_descartadas.docx`` and ``*_no_validas.xlsx`` when needed).

        Args:
            responses: The answers. Accepted: a single string; a list of strings; a list of
                ``(tema, texto)`` tuples; a dict ``{tema: texto}``; a list of dicts or a DataFrame
                with the columns 'Tema' and 'Respuesta'.
            output_filename (Optional[str]): Base name of the output files.
            output_dir (Optional[str]): Output folder (current folder if None).
            generate_docx (bool): Also write the DOCX for review.
            existing_bank_path (Optional[str]): Bank for the similarity filter.
            similarity_threshold (Optional[float]): Jaccard threshold (None disables the filter).
            default_topic (str): Topic for answers without one.

        Returns:
            pd.DataFrame: The accepted questions.
        """
        output_filename = output_filename or 'preguntas'
        output_dir = output_dir or os.getcwd()

        self.final_df = pd.DataFrame()
        self.rejected_questions_list = []
        self.invalid_questions_list = []
        self.duplicate_questions_count = 0
        self._set_similarity_threshold(similarity_threshold)
        self._load_existing_bank(existing_bank_path)

        accepted: List[Dict] = []
        seen = set()
        for topic, text in self._normalize_responses(responses, default_topic):
            accepted_in_response = 0
            for question in self._analyze_response(text, topic):
                key = (question["Tema"], question_identity(question))
                if key in seen:
                    self.duplicate_questions_count += 1
                    continue
                seen.add(key)
                accepted.append(question)
                accepted_in_response += 1
            self._log(f"Tema '{topic}': {accepted_in_response} preguntas aceptadas de esta respuesta.")
        if self.duplicate_questions_count:
            self._log(f"Se descartaron {self.duplicate_questions_count} preguntas repetidas (mismo tema, "
                      "enunciado y opciones).")

        self.final_df = pd.DataFrame(accepted, columns=QUESTION_COLUMNS) if accepted else pd.DataFrame()
        self._save_generation_outputs(output_filename, output_dir, generate_docx)
        self._log(
            f"Importación completada: {len(self.final_df)} aceptadas, "
            f"{len(self.rejected_questions_list)} descartadas por similitud, "
            f"{len(self.invalid_questions_list)} no válidas."
        )
        return self.final_df

    def import_questions_from_prompts_file(self, prompts_path: str, **kwargs) -> pd.DataFrame:
        """
        Imports the answers written in the 'Respuesta' column of a ``*_prompts.xlsx`` created
        by :meth:`export_prompts`. A cell may contain the answer itself or the path to a
        ``.txt``/``.md``/``.json`` file with it.

        Args:
            prompts_path (str): Path of the prompts workbook.
            **kwargs: Options of :meth:`import_questions_from_responses`.
        """
        df = pd.read_excel(prompts_path, sheet_name=0, dtype=str)
        if "Respuesta" not in df.columns:
            raise ValueError(f"El archivo '{prompts_path}' no tiene una columna 'Respuesta'.")

        base_dir = os.path.dirname(os.path.abspath(prompts_path))
        prompts_folder = os.path.splitext(os.path.abspath(prompts_path))[0]  # <nombre>_prompts/
        responses = []
        for _, row in df.iterrows():
            text = row.get("Respuesta")
            if not isinstance(text, str) or not text.strip():
                continue
            candidate = text.strip().strip('"').strip("'").strip()  # Windows "Copiar como ruta" adds quotes.
            if len(candidate) < 400 and "\n" not in candidate:
                candidates = [candidate] if os.path.isabs(candidate) else [os.path.join(base_dir, candidate)]
                candidates.append(os.path.join(prompts_folder, os.path.basename(candidate)))
                existing = next((path for path in candidates if os.path.isfile(path)), None)
                if existing:
                    with open(existing, "r", encoding="utf-8-sig") as f:
                        text = f.read()
            responses.append({"Tema": row.get("Tema") or "", "Respuesta": text})
        return self.import_questions_from_responses(responses, **kwargs)

    # ------------------------------------------------------------------
    # Generation through the API
    # ------------------------------------------------------------------
    def generate_multiple_choice_questions(
            self,
            pdf_paths: List[str],
            prompt_type: str,
            num_questions_per_chunk_target: int = 5,
            output_filename: Optional[str] = None,
            output_dir: Optional[str] = None,
            custom_prompt: Optional[str] = None,
            generate_docx: bool = True,
            existing_bank_path: Optional[str] = None,
            process_by_pages: bool = False,
            pages_per_chunk: int = 1,
            total_chunks: Optional[int] = None,
            chunking_mode: str = "pages",
            similarity_threshold: Optional[float] = 0.8,
            bank_prompt_scope: Optional[str] = None,
            prompt_example_content_type: str = "solo_enunciados",
            print_raw_gemini_answer: bool = False,
            max_generation_attempts_per_chunk: int = 3,
            include_similar_questions_in_prompt: bool = True,
            input_mode: str = "text",
            include_full_document_context: bool = True,
            progress_callback: Optional[Callable[[str], None]] = None,
            cancel_event: Optional[Any] = None,
    ) -> pd.DataFrame:
        """
        Main function to generate multiple-choice questions from multiple PDFs through the
        API of the configured provider (Gemini, OpenAI or Anthropic), with retries to reach
        the desired number of questions per chunk.

        This method orchestrates the entire generation process, including:
        - Loading an existing question bank for similarity checks.
        - Iterating through each provided PDF.
        - Splitting PDFs into chunks if required.
        - Looping with multiple attempts per chunk to reach the target number of questions.
        - Building prompts, generating content, and analyzing responses.
        - Saving accepted and rejected questions to files.

        Args:
            pdf_paths (List[str]): List of paths to the PDF files (or images).
            prompt_type (str): The key for the predefined prompt type.
            num_questions_per_chunk_target (int): The target number of questions per text chunk.
            output_filename (Optional[str]): The base name for the output files.
            output_dir (Optional[str]): The directory to save the output files.
                If None, uses the current working directory.
            custom_prompt (Optional[str]): A user-provided custom prompt to override the default.
            generate_docx (bool): Whether to generate a DOCX file for review.
            existing_bank_path (Optional[str]): Path to an existing Excel question bank for similarity filtering.
            process_by_pages (bool): If True, process the PDF in chunks of pages.
            pages_per_chunk (int): The number of pages per chunk if processing by pages.
            total_chunks (Optional[int]): Total chunks to split each PDF into.
            chunking_mode (str): Chunking strategy: 'pages' or 'text_length' (text mode only).
            similarity_threshold (Optional[float]): The threshold for the Jaccard similarity filter.
            bank_prompt_scope (Optional[str]): Scope for selecting guidance questions from the bank.
            prompt_example_content_type (str): Content type for guidance questions ('solo_enunciados' or 'enunciados_y_respuestas').
            print_raw_gemini_answer (bool): If True, prints the raw model answer for debugging.
            max_generation_attempts_per_chunk (int): The maximum number of attempts to reach the target per chunk.
            include_similar_questions_in_prompt (bool): If False, skips adding existing-bank questions to prompt.
            input_mode (str): 'text' (PDF text extraction), 'image' (images/PDF pages as images) or
                'pdf' (the PDF, or each PDF fragment, is attached as a document).
            include_full_document_context (bool): In 'text' mode with fragments, also send the whole
                document as context.
            progress_callback (Optional[Callable[[str], None]]): Receives progress messages.
            cancel_event (Optional[threading.Event]): When set, the generation stops after the current
                request and the questions obtained so far are saved (``self.cancelled`` becomes True).

        Returns:
            pd.DataFrame: A DataFrame containing all the successfully generated and filtered questions.

        Raises:
            QuotaExceededError, ServiceOverloadedError, InvalidAPIKeyError, ModelNotFoundError,
            ProviderNotAvailableError: The questions generated before the error are saved first.
        """
        self._validate_modes(input_mode, chunking_mode)
        # Missing or unsupported files are reported before the first (paid) API call.
        self._validate_sources(pdf_paths, input_mode)
        if not output_filename:
            output_filename = 'preguntas'

        if output_dir is None:
            output_dir = os.getcwd()

        os.makedirs(output_dir, exist_ok=True)

        self._progress_callback = progress_callback
        self.final_df = pd.DataFrame()
        self.rejected_questions_list = []  # Accumulates all rejected questions
        self.invalid_questions_list = []
        self.cancelled = False
        self.duplicate_questions_count = 0
        self.last_provider_error = None

        self._set_similarity_threshold(similarity_threshold)
        self.current_prompt_example_content_type = prompt_example_content_type
        self._log(f"Contenido de ejemplos en prompt: {self.current_prompt_example_content_type}")
        self._load_existing_bank(existing_bank_path)

        guidance_df = self.existing_bank_df if include_similar_questions_in_prompt else None
        guidance_scope = bank_prompt_scope if include_similar_questions_in_prompt else None
        all_accepted: List[Dict] = []
        accepted_keys = set()

        def is_cancelled() -> bool:
            if cancel_event is not None and cancel_event.is_set():
                self.cancelled = True
            return self.cancelled

        try:
            for pdf_path in pdf_paths:
                if is_cancelled():
                    break
                self._log(f"\n--- Procesando archivo fuente: {pdf_path} ---")
                current_pdf_topic = self.extract_topic_number_from_path(pdf_path)

                if input_mode == "text":
                    chunks, full_text = self._plan_text_chunks(
                        pdf_path, process_by_pages, pages_per_chunk, total_chunks, chunking_mode)
                    use_context = process_by_pages and include_full_document_context and len(chunks) > 1
                else:
                    chunks = self._plan_attachment_chunks(
                        pdf_path, input_mode, process_by_pages, pages_per_chunk, total_chunks, load_bytes=True)
                    full_text, use_context = None, False
                if not chunks:
                    continue

                for i, chunk in enumerate(chunks):
                    if is_cancelled():
                        break
                    if input_mode == "text":
                        material = chunk["text"]
                    else:
                        # Page images are not named in the prompt (there may be dozens of them).
                        names = [item.get("name", "") for item in chunk["attachments"]
                                 if item.get("name") and input_mode == "pdf"]
                        material = self._attachment_instruction(input_mode, names, chunk["label"])

                    self._log(f"\nProcesando fragmento {i + 1}/{len(chunks)} ({chunk['label']})...")
                    questions_accepted_for_this_chunk: List[Dict] = []
                    attempts_for_this_chunk = 0

                    while len(questions_accepted_for_this_chunk) < num_questions_per_chunk_target and \
                            attempts_for_this_chunk < max_generation_attempts_per_chunk:
                        if is_cancelled():
                            break
                        attempts_for_this_chunk += 1
                        num_still_needed = num_questions_per_chunk_target - len(questions_accepted_for_this_chunk)
                        self._log(f"  Intento {attempts_for_this_chunk}/{max_generation_attempts_per_chunk} para el fragmento: Se necesitan {num_still_needed} preguntas más.")

                        prompt = self._build_prompt(
                            working_chunk_text=material,
                            num_questions_to_generate=num_still_needed,
                            prompt_type=prompt_type,
                            custom_prompt=custom_prompt,
                            existing_questions_df=guidance_df,
                            current_topic=current_pdf_topic,
                            full_document_context_text=full_text if use_context else None,
                            bank_prompt_scope=guidance_scope,
                            questions_from_current_chunk_attempts=questions_accepted_for_this_chunk,
                            material_is_attachment=input_mode != "text",
                        )

                        try:
                            response_text = self.ai_client.generate(prompt, chunk["attachments"])
                        except _FATAL_PROVIDER_ERRORS:
                            raise
                        except Exception as e_gen:
                            # Other errors are logged and the next attempt is tried.
                            self.last_provider_error = str(e_gen)
                            self._log(f"    Error durante la generación en el intento {attempts_for_this_chunk} para el fragmento {i + 1}: {e_gen}")
                            if not isinstance(e_gen, AIProviderError):
                                print(traceback.format_exc())
                            continue

                        if print_raw_gemini_answer:
                            print(f"    Respuesta cruda del modelo (Fragmento {i + 1}, Intento {attempts_for_this_chunk}):\n    >>>\n{response_text}\n    <<<\n    ---")

                        time.sleep(1)
                        unique_newly_accepted = []
                        for q_new in self._analyze_response(response_text, current_pdf_topic):
                            # Discard repetitions (same topic, statement and options) within the run.
                            key = (q_new["Tema"], question_identity(q_new))
                            if key in accepted_keys:
                                self.duplicate_questions_count += 1
                                continue
                            accepted_keys.add(key)
                            unique_newly_accepted.append(q_new)

                        questions_accepted_for_this_chunk.extend(unique_newly_accepted)
                        # Kept at once, so an error in a later attempt does not lose them.
                        all_accepted.extend(unique_newly_accepted)
                        self._log(
                            f"    Intento {attempts_for_this_chunk}: {len(unique_newly_accepted)} preguntas nuevas aceptadas para este fragmento. Total para fragmento: {len(questions_accepted_for_this_chunk)}/{num_questions_per_chunk_target}")

                    if len(questions_accepted_for_this_chunk) < num_questions_per_chunk_target and not self.cancelled:
                        self._log(
                            f"  Advertencia: No se alcanzó el objetivo de {num_questions_per_chunk_target} preguntas para el fragmento {i + 1} después de {max_generation_attempts_per_chunk} intentos. Se obtuvieron {len(questions_accepted_for_this_chunk)}.")
        except Exception:
            # Fatal provider errors and unexpected ones (unreadable PDF, file locked...): the questions
            # obtained so far are saved before the error reaches the caller.
            if all_accepted:
                self._log(f"Guardando las {len(all_accepted)} preguntas obtenidas antes del error...")
                self.final_df = pd.DataFrame(all_accepted, columns=QUESTION_COLUMNS)
                try:
                    self._save_generation_outputs(output_filename, output_dir, generate_docx)
                except Exception as save_error:
                    self._log(f"No se pudieron guardar las preguntas obtenidas: {save_error}")
            self._progress_callback = None
            raise

        if self.cancelled:
            self._log("Generación cancelada por el usuario: se guardan las preguntas obtenidas hasta ahora.")

        self.final_df = pd.DataFrame(all_accepted, columns=QUESTION_COLUMNS) if all_accepted else pd.DataFrame()
        self._save_generation_outputs(output_filename, output_dir, generate_docx)

        self._log("\n--- Proceso de generación completado. ---")
        self._log(f"Total preguntas aceptadas y guardadas: {len(self.final_df)}")
        if self.rejected_questions_list:
            self._log(f"Total preguntas rechazadas por similitud (acumulado de todos los intentos): {len(self.rejected_questions_list)}")
        if self.invalid_questions_list:
            self._log(f"Total fragmentos de respuesta no válidos (formato): {len(self.invalid_questions_list)}")
        if self.duplicate_questions_count:
            self._log(f"Total preguntas repetidas descartadas: {self.duplicate_questions_count}")
        self._progress_callback = None

        return self.final_df

# Example of use (uncomment to test)
# if __name__ == '__main__':
#     # With the API (provider: 'gemini', 'openai' or 'anthropic'):
#     generator = QuestionGenerator(api_key="YOUR_API_KEY", model_name="gemini-flash-latest", provider="gemini")
#     questions_df = generator.generate_multiple_choice_questions(
#         ['TEMA 05_Inspecciones de seguridad.pdf'], prompt_type="PRL",
#         num_questions_per_chunk_target=10, output_filename="banco_de_preguntas")
#
#     # Without the API (Google AI Studio, ChatGPT or Claude in the browser):
#     generator = QuestionGenerator()
#     prompts_df = generator.build_generation_prompts(['TEMA 05_Inspecciones de seguridad.pdf'], prompt_type="PRL")
#     generator.export_prompts(prompts_df, output_filename="banco_de_preguntas")
#     # ... paste each prompt in the chat and the answers in the 'Respuesta' column, then:
#     generator.import_questions_from_prompts_file("banco_de_preguntas_prompts.xlsx",
#                                                  output_filename="banco_de_preguntas")
